from __future__ import annotations

import json
import os
import shutil
import socket
import tempfile
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

SCHEMA: dict[str, list[str]] = {
    "Settings": ["Key", "Value", "Description"],
    "Machine_Master": [
        "Machine_ID",
        "Machine_Name",
        "Machine_Type",
        "Capacity",
        "Prize_Category",
        "Active",
        "Sort_Order",
        "Notes",
        "Created_At",
        "Updated_At",
        "Barcodes_JSON",
        "Image_File",
    ],
    "Daily_Closing": [
        "Closing_ID",
        "Report_Date",
        "Outlet",
        "Cash_Sales",
        "Cash_Transactions",
        "ABA_Sales",
        "ABA_Transactions",
        "Adjustment",
        "Beginning_Coins",
        "Coins_Added",
        "Final_Coins",
        "Coins_Dispensed",
        "Total_Sales",
        "Total_Transactions",
        "Average_Sale_Value_Per_Coin",
        "Machine_Coins_Used",
        "Coin_Variance",
        "Total_Prizes_Won",
        "Average_Coins_Per_Prize",
        "Average_Revenue_Per_Prize",
        "Closing_Status",
        "Workflow_Status",
        "Closed_By",
        "Verified_By",
        "Notes",
        "Created_At",
        "Updated_At",
        "Finalized_At",
        "Cash_Sales_KHR",
        "Exchange_Rate_USD_KHR",
        "Price_Per_Coin_USD",
        "Coins_Used",
        "Coin_Return",
        "Lose_Over",
        "Discount_USD",
        "Discount_Percent",
        "Overall_Win_Rate",
    ],
    "Machine_Closing": [
        "Closing_ID",
        "Report_Date",
        "Machine_ID",
        "Machine_Name",
        "Machine_Type",
        "Capacity",
        "Begin_Prize",
        "Refill_Prize",
        "Available_Prize",
        "Final_Prize",
        "Prizes_Won",
        "Refill_Needed",
        "Begin_Coin_Meter",
        "Final_Coin_Meter",
        "Coins_Used",
        "Coins_Per_Prize",
        "Allocated_Sales",
        "Average_Revenue_Per_Prize",
        "Machine_Status",
        "Notes",
        "Created_At",
        "Updated_At",
    ],
    "Audit_Log": [
        "Audit_ID",
        "Timestamp",
        "Action",
        "Entity_Type",
        "Entity_ID",
        "User_Name",
        "Computer_Name",
        "Details_JSON",
    ],
}


class DatabaseLockedError(RuntimeError):
    pass


class ExcelDatabase:
    def __init__(self, path: Path, backup_dir: Path) -> None:
        self.path = Path(path)
        self.backup_dir = Path(backup_dir)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        self.lock_path = self.path.with_suffix(self.path.suffix + ".lock")
        self._lock_fd: int | None = None
        self._cache_mtime_ns: int = -1
        self._cache_settings: dict[str, Any] = {}
        self._cache_machines: list[dict[str, Any]] = []
        self._cache_closings: list[dict[str, Any]] = []
        self._cache_machine_closings: list[dict[str, Any]] = []
        if not self.path.exists():
            self.create_default_database(self.path)
        else:
            self._migrate_schema()
        self._validate_schema()

    @staticmethod
    def _process_is_running(pid: int) -> bool:
        if pid <= 0:
            return False
        if os.name == "nt":
            import ctypes

            PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
            SYNCHRONIZE = 0x00100000
            WAIT_TIMEOUT = 0x00000102
            kernel32 = ctypes.windll.kernel32
            kernel32.OpenProcess.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
            kernel32.OpenProcess.restype = ctypes.c_void_p
            kernel32.WaitForSingleObject.argtypes = [ctypes.c_void_p, ctypes.c_uint32]
            kernel32.WaitForSingleObject.restype = ctypes.c_uint32
            kernel32.CloseHandle.argtypes = [ctypes.c_void_p]
            kernel32.CloseHandle.restype = ctypes.c_int
            handle = kernel32.OpenProcess(
                PROCESS_QUERY_LIMITED_INFORMATION | SYNCHRONIZE,
                False,
                int(pid),
            )
            if not handle:
                return False
            try:
                return kernel32.WaitForSingleObject(handle, 0) == WAIT_TIMEOUT
            finally:
                kernel32.CloseHandle(handle)
        try:
            os.kill(pid, 0)
        except ProcessLookupError:
            return False
        except PermissionError:
            return True
        except OSError:
            return False
        return True

    def _remove_stale_lock(self) -> None:
        if not self.lock_path.exists():
            return
        try:
            details = json.loads(self.lock_path.read_text(encoding="utf-8"))
            pid = int(details.get("pid") or 0)
            computer = str(details.get("computer") or "")
        except (OSError, ValueError, TypeError, json.JSONDecodeError):
            pid = 0
            computer = ""

        same_computer = not computer or computer.casefold() == socket.gethostname().casefold()
        if same_computer and not self._process_is_running(pid):
            try:
                self.lock_path.unlink(missing_ok=True)
            except OSError:
                pass

    def acquire_application_lock(self) -> None:
        self._remove_stale_lock()
        try:
            self._lock_fd = os.open(self.lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            payload = {
                "pid": os.getpid(),
                "computer": socket.gethostname(),
                "started_at": datetime.now().isoformat(timespec="seconds"),
            }
            os.write(self._lock_fd, json.dumps(payload).encode("utf-8"))
        except FileExistsError as exc:
            details = ""
            try:
                details = self.lock_path.read_text(encoding="utf-8")
            except OSError:
                pass
            raise DatabaseLockedError(
                "The Excel database is already open by another application session.\n"
                f"Lock details: {details or 'Unavailable'}"
            ) from exc

    def release_application_lock(self) -> None:
        if self._lock_fd is not None:
            try:
                os.close(self._lock_fd)
            except OSError:
                pass
            self._lock_fd = None
        try:
            self.lock_path.unlink(missing_ok=True)
        except OSError:
            pass

    @staticmethod
    def create_default_database(path: Path) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, headers in SCHEMA.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="102A43")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"

        settings = workbook["Settings"]
        settings_rows = [
            ["App_Name", "JOLI POLI Claw Closing", "Application title"],
            ["Outlet", "JOLI POLI", "Default outlet"],
            ["Currency", "USD", "Report currency"],
            ["Variance_Tolerance", 0, "Allowed coin difference before variance status"],
            ["Database_Version", "1.2", "Excel database schema version"],
        ]
        for row in settings_rows:
            settings.append(row)

        machine_sheet = workbook["Machine_Master"]
        now = datetime.now().isoformat(timespec="seconds")
        machines: list[list[Any]] = []
        order = 1
        for prefix, machine_type, count, capacity, category in [
            ("CL", "Claw / Plush", 7, 40, "Plush Toy"),
            ("KC", "Keychain", 7, 20, "Keychain"),
            ("BB", "Blind Box", 2, 50, "Blind Box"),
        ]:
            for index in range(1, count + 1):
                machine_id = f"{prefix}-{index:02d}"
                machines.append(
                    [
                        machine_id,
                        f"{machine_type} {index:02d}",
                        machine_type,
                        capacity,
                        category,
                        True,
                        order,
                        "Imported from the current closing workbook structure",
                        now,
                        now,
                        "[]",
                        "",
                    ]
                )
                order += 1
        for row in machines:
            machine_sheet.append(row)

        widths = {
            "Settings": [24, 24, 48],
            "Machine_Master": [14, 24, 18, 12, 18, 10, 12, 48, 22, 22, 44, 34],
            "Daily_Closing": [18] * len(SCHEMA["Daily_Closing"]),
            "Machine_Closing": [18] * len(SCHEMA["Machine_Closing"]),
            "Audit_Log": [18, 22, 18, 18, 22, 18, 20, 55],
        }
        for sheet_name, sheet_widths in widths.items():
            sheet = workbook[sheet_name]
            for idx, width in enumerate(sheet_widths, start=1):
                sheet.column_dimensions[sheet.cell(1, idx).column_letter].width = width

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)

    def _migrate_schema(self) -> None:
        """Append newly introduced columns without changing existing workbook data."""
        workbook = load_workbook(self.path)
        changed = False
        try:
            for sheet_name, expected_headers in SCHEMA.items():
                if sheet_name not in workbook.sheetnames:
                    continue
                sheet = workbook[sheet_name]
                actual_headers = [
                    str(cell.value) if cell.value is not None else ""
                    for cell in sheet[1]
                ]
                missing = [header for header in expected_headers if header not in actual_headers]
                if not missing:
                    continue

                existing_expected = [
                    header for header in expected_headers if header in actual_headers
                ]
                if actual_headers[: len(existing_expected)] != existing_expected:
                    raise RuntimeError(
                        f"Database sheet {sheet_name!r} cannot be migrated safely."
                    )

                for header in missing:
                    column = sheet.max_column + 1
                    cell = sheet.cell(1, column)
                    cell.value = header
                    cell.fill = PatternFill("solid", fgColor="102A43")
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    sheet.column_dimensions[cell.column_letter].width = (
                        44 if header == "Barcodes_JSON" else 34
                    )
                    changed = True

                sheet.auto_filter.ref = f"A1:{sheet.cell(1, sheet.max_column).coordinate}"

            if "Settings" in workbook.sheetnames:
                settings = workbook["Settings"]
                headers = self._header_map(settings)
                database_version_row = None
                for row_idx in range(2, self._max_row(settings) + 1):
                    if str(settings.cell(row_idx, headers["Key"]).value or "") == "Database_Version":
                        database_version_row = row_idx
                        break
                if database_version_row is None:
                    database_version_row = self._max_row(settings) + 1
                    settings.cell(database_version_row, headers["Key"]).value = "Database_Version"
                    settings.cell(database_version_row, headers["Description"]).value = (
                        "Excel database schema version"
                    )
                    changed = True
                current_version = str(settings.cell(database_version_row, headers["Value"]).value or "").strip()
                # A later v2.1.78 extension owns newer schema versions.  The
                # base migration may initialise an empty/older workbook but
                # must never downgrade it, which would force a needless save.
                if not current_version or current_version < "1.2":
                    settings.cell(database_version_row, headers["Value"]).value = "1.2"
                    changed = True

            if changed:
                self._save_atomic(workbook)
        finally:
            workbook.close()

    def _validate_schema(self) -> None:
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            missing_sheets = [name for name in SCHEMA if name not in workbook.sheetnames]
            if missing_sheets:
                raise RuntimeError(f"Database is missing required sheets: {', '.join(missing_sheets)}")
            for name, expected_headers in SCHEMA.items():
                sheet = workbook[name]
                actual = [sheet.cell(1, idx).value for idx in range(1, len(expected_headers) + 1)]
                if actual != expected_headers:
                    raise RuntimeError(f"Database sheet {name!r} has an invalid header structure.")
        finally:
            workbook.close()

    @contextmanager
    def _workbook(self, read_only: bool = False):
        workbook = load_workbook(self.path, read_only=read_only, data_only=False)
        try:
            yield workbook
        finally:
            workbook.close()


    @staticmethod
    def _max_row(sheet) -> int:
        value = sheet.max_row
        if value is not None:
            return int(value)
        dimension = sheet.calculate_dimension(force=True)
        end_ref = dimension.split(":")[-1]
        digits = "".join(character for character in end_ref if character.isdigit())
        return int(digits or 1)

    @staticmethod
    def _header_map(sheet) -> dict[str, int]:
        return {cell.value: cell.column for cell in sheet[1] if cell.value}

    @staticmethod
    def _row_to_dict(sheet, row_index: int) -> dict[str, Any]:
        headers = [cell.value for cell in sheet[1]]
        return {headers[idx]: sheet.cell(row_index, idx + 1).value for idx in range(len(headers))}

    def _invalidate_cache(self) -> None:
        self._cache_mtime_ns = -1
        self._cache_settings = {}
        self._cache_machines = []
        self._cache_closings = []
        self._cache_machine_closings = []

    @staticmethod
    def _sheet_rows(sheet) -> list[dict[str, Any]]:
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(iterator))
        except StopIteration:
            return []
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if not any(value is not None for value in values):
                continue
            rows.append({headers[index]: values[index] for index in range(len(headers))})
        return rows

    def _ensure_cache(self) -> None:
        try:
            mtime_ns = self.path.stat().st_mtime_ns
        except OSError:
            mtime_ns = -1
        if self._cache_mtime_ns == mtime_ns and self._cache_settings:
            return

        with self._workbook(read_only=True) as workbook:
            settings_rows = self._sheet_rows(workbook["Settings"])
            machines = self._sheet_rows(workbook["Machine_Master"])
            closings = self._sheet_rows(workbook["Daily_Closing"])
            machine_closings = self._sheet_rows(workbook["Machine_Closing"])

        self._cache_settings = {
            str(row.get("Key")): row.get("Value")
            for row in settings_rows
            if row.get("Key") is not None
        }
        self._cache_machines = machines
        self._cache_closings = closings
        self._cache_machine_closings = machine_closings
        self._cache_mtime_ns = mtime_ns

    def _backup(self) -> None:
        if not self.path.exists():
            return
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        destination = self.backup_dir / f"claw_machine_database-{stamp}.xlsx"
        shutil.copy2(self.path, destination)
        backups = sorted(self.backup_dir.glob("claw_machine_database-*.xlsx"), reverse=True)
        for old in backups[40:]:
            old.unlink(missing_ok=True)

    def _save_atomic(self, workbook) -> None:
        self._backup()
        with tempfile.NamedTemporaryFile(
            prefix="claw-db-", suffix=".xlsx", dir=self.path.parent, delete=False
        ) as handle:
            temp_path = Path(handle.name)
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self.path)
            self._invalidate_cache()
        finally:
            temp_path.unlink(missing_ok=True)

    def get_setting(self, key: str, default: Any = None) -> Any:
        self._ensure_cache()
        return self._cache_settings.get(key, default)

    def set_settings(self, values: dict[str, tuple[Any, str]]) -> dict[str, Any]:
        """Save application settings atomically without changing historical business records."""
        workbook = load_workbook(self.path)
        try:
            sheet = workbook["Settings"]
            headers = self._header_map(sheet)
            rows_by_key: dict[str, int] = {}
            for row_idx in range(2, self._max_row(sheet) + 1):
                key = str(sheet.cell(row_idx, headers["Key"]).value or "").strip()
                if key:
                    rows_by_key[key] = row_idx

            saved: dict[str, Any] = {}
            for key, (value, description) in values.items():
                row_idx = rows_by_key.get(key)
                if row_idx is None:
                    row_idx = self._max_row(sheet) + 1
                    rows_by_key[key] = row_idx
                    sheet.cell(row_idx, headers["Key"]).value = key
                sheet.cell(row_idx, headers["Value"]).value = value
                sheet.cell(row_idx, headers["Description"]).value = description
                saved[key] = value

            self._save_atomic(workbook)
            return saved
        finally:
            workbook.close()

    def migrate_branding(self, app_name: str, outlet: str) -> bool:
        """Update legacy Hygge Zone settings without changing historical closings."""
        workbook = load_workbook(self.path)
        changed = False
        try:
            sheet = workbook["Settings"]
            headers = self._header_map(sheet)
            rows_by_key: dict[str, int] = {}
            for row_idx in range(2, self._max_row(sheet) + 1):
                key = str(sheet.cell(row_idx, headers["Key"]).value or "").strip()
                if key:
                    rows_by_key[key] = row_idx

            desired = {
                "App_Name": (app_name, "Application title"),
                "Outlet": (outlet, "Default outlet"),
            }
            for key, (value, description) in desired.items():
                row_idx = rows_by_key.get(key)
                if row_idx is None:
                    row_idx = self._max_row(sheet) + 1
                    sheet.cell(row_idx, headers["Key"]).value = key
                    sheet.cell(row_idx, headers["Value"]).value = value
                    sheet.cell(row_idx, headers["Description"]).value = description
                    changed = True
                    continue

                current = str(sheet.cell(row_idx, headers["Value"]).value or "").strip()
                should_replace = key == "App_Name" or not current or current.casefold() == "hygge zone"
                if should_replace and current != value:
                    sheet.cell(row_idx, headers["Value"]).value = value
                    changed = True
                if not sheet.cell(row_idx, headers["Description"]).value:
                    sheet.cell(row_idx, headers["Description"]).value = description
                    changed = True

            if changed:
                self._save_atomic(workbook)
            return changed
        finally:
            workbook.close()

    def list_machines(self, active_only: bool = True) -> list[dict[str, Any]]:
        self._ensure_cache()
        rows = [dict(row) for row in self._cache_machines]
        if active_only:
            rows = [row for row in rows if bool(row.get("Active"))]
        type_order = {"Claw / Plush": 0, "Keychain": 1, "Blind Box": 2}

        def machine_number(row: dict[str, Any]) -> int:
            machine_id = str(row.get("Machine_ID") or "")
            digits = "".join(character for character in machine_id if character.isdigit())
            return int(digits) if digits else 9999

        rows.sort(
            key=lambda row: (
                type_order.get(str(row.get("Machine_Type") or ""), 99),
                machine_number(row),
                str(row.get("Machine_ID") or ""),
            )
        )
        return rows

    def upsert_machine(self, machine: dict[str, Any], user_name: str) -> None:
        workbook = load_workbook(self.path)
        try:
            sheet = workbook["Machine_Master"]
            headers = self._header_map(sheet)
            target_row = None
            for row_idx in range(2, self._max_row(sheet) + 1):
                if sheet.cell(row_idx, headers["Machine_ID"]).value == machine["Machine_ID"]:
                    target_row = row_idx
                    break
            now = datetime.now().isoformat(timespec="seconds")
            if target_row is None:
                target_row = self._max_row(sheet) + 1
                machine["Created_At"] = now
            else:
                machine["Created_At"] = sheet.cell(target_row, headers["Created_At"]).value or now
            machine["Updated_At"] = now
            for header in SCHEMA["Machine_Master"]:
                sheet.cell(target_row, headers[header]).value = machine.get(header)
            self._append_audit(
                workbook,
                action="UPSERT",
                entity_type="Machine",
                entity_id=str(machine["Machine_ID"]),
                user_name=user_name,
                details=machine,
            )
            self._save_atomic(workbook)
        finally:
            workbook.close()

    def next_closing_id(self, report_date: date) -> str:
        self._ensure_cache()
        prefix = f"CL-{report_date.strftime('%Y%m%d')}-"
        maximum = 0
        for row in self._cache_closings:
            value = str(row.get("Closing_ID") or "")
            if value.startswith(prefix):
                try:
                    maximum = max(maximum, int(value.rsplit("-", 1)[1]))
                except ValueError:
                    continue
        return f"{prefix}{maximum + 1:03d}"

    def save_closing(
        self,
        header: dict[str, Any],
        machine_rows: list[dict[str, Any]],
        user_name: str,
    ) -> str:
        workbook = load_workbook(self.path)
        try:
            now = datetime.now().isoformat(timespec="seconds")
            daily = workbook["Daily_Closing"]
            daily_headers = self._header_map(daily)
            closing_id = str(header.get("Closing_ID") or "")
            target_row = None
            for row_idx in range(2, self._max_row(daily) + 1):
                if str(daily.cell(row_idx, daily_headers["Closing_ID"]).value or "") == closing_id:
                    target_row = row_idx
                    break
            if target_row is None:
                target_row = self._max_row(daily) + 1
                header["Created_At"] = now
            else:
                existing_status = daily.cell(target_row, daily_headers["Workflow_Status"]).value
                if existing_status == "Finalized":
                    raise ValueError("A finalized closing cannot be overwritten.")
                header["Created_At"] = daily.cell(target_row, daily_headers["Created_At"]).value or now
            header["Updated_At"] = now
            if header.get("Workflow_Status") == "Finalized":
                header["Finalized_At"] = now

            for column_name in SCHEMA["Daily_Closing"]:
                daily.cell(target_row, daily_headers[column_name]).value = header.get(column_name)

            machine_sheet = workbook["Machine_Closing"]
            machine_headers = self._header_map(machine_sheet)
            rows_to_delete: list[int] = []
            for row_idx in range(2, self._max_row(machine_sheet) + 1):
                if str(machine_sheet.cell(row_idx, machine_headers["Closing_ID"]).value or "") == closing_id:
                    rows_to_delete.append(row_idx)
            for row_idx in reversed(rows_to_delete):
                machine_sheet.delete_rows(row_idx, 1)

            for machine in machine_rows:
                machine["Closing_ID"] = closing_id
                machine["Report_Date"] = header.get("Report_Date")
                machine["Created_At"] = now
                machine["Updated_At"] = now
                row_index = self._max_row(machine_sheet) + 1
                for column_name in SCHEMA["Machine_Closing"]:
                    machine_sheet.cell(row_index, machine_headers[column_name]).value = machine.get(column_name)

            self._append_audit(
                workbook,
                action="FINALIZE" if header.get("Workflow_Status") == "Finalized" else "SAVE_DRAFT",
                entity_type="Closing",
                entity_id=closing_id,
                user_name=user_name,
                details={
                    "report_date": str(header.get("Report_Date")),
                    "workflow_status": header.get("Workflow_Status"),
                    "machine_count": len(machine_rows),
                },
            )
            self._save_atomic(workbook)
            return closing_id
        finally:
            workbook.close()

    def _append_audit(
        self,
        workbook,
        action: str,
        entity_type: str,
        entity_id: str,
        user_name: str,
        details: dict[str, Any],
    ) -> None:
        sheet = workbook["Audit_Log"]
        audit_id = f"AUD-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
        sheet.append(
            [
                audit_id,
                datetime.now().isoformat(timespec="seconds"),
                action,
                entity_type,
                entity_id,
                user_name,
                socket.gethostname(),
                json.dumps(details, ensure_ascii=False, default=str),
            ]
        )

    def list_closings(self, limit: int = 500) -> list[dict[str, Any]]:
        self._ensure_cache()
        rows = [dict(row) for row in self._cache_closings]
        rows.sort(
            key=lambda row: (str(row.get("Report_Date") or ""), str(row.get("Created_At") or "")),
            reverse=True,
        )
        return rows[:limit]

    def get_closing(self, closing_id: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        self._ensure_cache()
        header = next(
            (
                dict(row)
                for row in self._cache_closings
                if str(row.get("Closing_ID") or "") == closing_id
            ),
            None,
        )
        if header is None:
            raise KeyError(f"Closing {closing_id!r} was not found.")
        machines = [
            dict(row)
            for row in self._cache_machine_closings
            if str(row.get("Closing_ID") or "") == closing_id
        ]
        return header, machines

    def previous_machine_state(self, machine_id: str, before_date: date | None = None) -> dict[str, Any]:
        return self.previous_machine_states([machine_id], before_date).get(machine_id, {})

    def previous_machine_states(
        self,
        machine_ids: Iterable[str],
        before_date: date | None = None,
    ) -> dict[str, dict[str, Any]]:
        self._ensure_cache()
        requested = {str(machine_id) for machine_id in machine_ids}
        if not requested:
            return {}

        finalized: dict[str, tuple[str, str]] = {}
        for row in self._cache_closings:
            if row.get("Workflow_Status") != "Finalized":
                continue
            closing_id = str(row.get("Closing_ID") or "")
            report_date = str(row.get("Report_Date") or "")[:10]
            if before_date and report_date >= before_date.isoformat():
                continue
            finalized[closing_id] = (
                report_date,
                str(row.get("Finalized_At") or ""),
            )

        latest: dict[str, tuple[tuple[str, str], dict[str, Any]]] = {}
        for row in self._cache_machine_closings:
            machine_id = str(row.get("Machine_ID") or "")
            if machine_id not in requested:
                continue
            closing_id = str(row.get("Closing_ID") or "")
            candidate_key = finalized.get(closing_id)
            if candidate_key is None:
                continue
            current = latest.get(machine_id)
            if current is None or candidate_key > current[0]:
                latest[machine_id] = (candidate_key, dict(row))

        return {machine_id: value[1] for machine_id, value in latest.items()}

    def dashboard_summary(self) -> dict[str, Any]:
        closings = [row for row in self.list_closings(1000) if row.get("Workflow_Status") == "Finalized"]
        latest = closings[0] if closings else None
        current_month = date.today().strftime("%Y-%m")
        month_rows = [row for row in closings if str(row.get("Report_Date") or "").startswith(current_month)]
        return {
            "latest": latest,
            "month_total_sales": sum(float(row.get("Total_Sales") or 0) for row in month_rows),
            "month_prizes": sum(int(row.get("Total_Prizes_Won") or 0) for row in month_rows),
            "month_coins": sum(int(row.get("Machine_Coins_Used") or 0) for row in month_rows),
            "month_closings": len(month_rows),
        }

    def monthly_closings(self, month_text: str) -> list[dict[str, Any]]:
        return [
            row
            for row in self.list_closings(5000)
            if row.get("Workflow_Status") == "Finalized"
            and str(row.get("Report_Date") or "").startswith(month_text)
        ]

    def monthly_machine_rows(self, month_text: str) -> list[dict[str, Any]]:
        valid_ids = {row["Closing_ID"] for row in self.monthly_closings(month_text)}
        if not valid_ids:
            return []
        self._ensure_cache()
        return [
            dict(row)
            for row in self._cache_machine_closings
            if row.get("Closing_ID") in valid_ids
        ]
