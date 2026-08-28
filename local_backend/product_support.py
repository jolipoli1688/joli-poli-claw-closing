from __future__ import annotations

import base64
import binascii
import hashlib
import json
import os
import re
import urllib.parse
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from fastapi import HTTPException, Query
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from pdf_report import build_daily_closing_pdf

PRODUCT_CLOSING_HEADERS = [
    "Closing_ID",
    "Report_Date",
    "Machine_ID",
    "Machine_Name",
    "Product_ID",
    "Barcode",
    "Image_File",
    "Begin_Qty",
    "Final_Qty",
    "Qty_Used",
    "Created_At",
    "Updated_At",
]

PRODUCT_IMAGE_MAX_BYTES = 5 * 1024 * 1024
IMAGE_MIME_EXTENSIONS = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
}


def _to_int(value: Any, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(str(value).replace(",", "").strip()))
    except (TypeError, ValueError):
        return default


def _safe_product_id(value: Any, machine_id: str, barcode: str, index: int) -> str:
    raw = re.sub(r"[^A-Za-z0-9_.-]+", "-", str(value or "")).strip("-")
    if raw:
        return raw[:80]
    digest = hashlib.sha1(f"{machine_id}|{barcode}|{index}".encode("utf-8")).hexdigest()[:12]
    return f"product-{digest}"


def _parse_json_list(value: Any) -> list[Any]:
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    text = str(value or "").strip()
    if not text:
        return []
    try:
        loaded = json.loads(text)
    except json.JSONDecodeError:
        return [part.strip() for part in re.split(r"[\r\n,;|]+", text) if part.strip()]
    return loaded if isinstance(loaded, list) else []


def normalise_products(
    value: Any,
    machine_id: str = "",
    fallback_barcodes: Any = None,
    fallback_image: Any = "",
) -> list[dict[str, Any]]:
    raw_values = _parse_json_list(value)
    if not raw_values:
        raw_values = _parse_json_list(fallback_barcodes)

    legacy_image = Path(str(fallback_image or "")).name
    if not raw_values and legacy_image:
        raw_values = [{}]
    products: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    seen_codes: set[str] = set()

    for index, item in enumerate(raw_values):
        source = item if isinstance(item, dict) else {"barcode": item}
        barcode = str(source.get("barcode") or source.get("code") or "").strip()
        default_quantity = 0
        image_file = Path(str(source.get("image_file") or "")).name
        if not image_file and index == 0 and legacy_image:
            image_file = legacy_image

        product_id = _safe_product_id(
            source.get("product_id") or source.get("id"),
            machine_id,
            barcode,
            index,
        )
        base_id = product_id
        suffix = 2
        while product_id.casefold() in seen_ids:
            product_id = f"{base_id}-{suffix}"
            suffix += 1
        seen_ids.add(product_id.casefold())

        if barcode:
            code_key = barcode.casefold()
            if code_key in seen_codes:
                continue
            seen_codes.add(code_key)

        if not barcode and not image_file and default_quantity == 0:
            continue
        products.append(
            {
                "product_id": product_id,
                "barcode": barcode,
                "default_quantity": default_quantity,
                "image_file": image_file,
            }
        )
        if len(products) > 30:
            raise ValueError("A machine can contain up to 30 products.")

    return products


def product_image_url(file_name: Any) -> str:
    safe_name = Path(str(file_name or "")).name
    return f"/api/machine-images/{urllib.parse.quote(safe_name)}" if safe_name else ""


def product_view(product: dict[str, Any]) -> dict[str, Any]:
    result = dict(product)
    result["image_url"] = product_image_url(result.get("image_file"))
    return result


def machine_product_view(record: dict[str, Any]) -> dict[str, Any]:
    result = dict(record)
    products = normalise_products(
        result.get("Products_JSON"),
        str(result.get("Machine_ID") or ""),
        result.get("Barcodes_JSON"),
        result.get("Image_File"),
    )
    result["Products"] = [product_view(item) for item in products]
    result["Barcodes"] = [item["barcode"] for item in products if item.get("barcode")]
    first_image = next((item.get("image_url") for item in result["Products"] if item.get("image_url")), "")
    result["Image_URL"] = first_image
    return result


def _style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="102A43")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, sheet.max_column).coordinate}"


def _upgrade_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    changed = False
    try:
        if "Product_Closing" not in workbook.sheetnames:
            sheet = workbook.create_sheet("Product_Closing")
            sheet.append(PRODUCT_CLOSING_HEADERS)
            _style_header(sheet)
            widths = [18, 14, 14, 24, 24, 24, 34, 14, 14, 14, 22, 22]
            for index, width in enumerate(widths, start=1):
                sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
            changed = True

        machine_sheet = workbook["Machine_Master"]
        headers = {cell.value: cell.column for cell in machine_sheet[1] if cell.value}
        if "Products_JSON" not in headers:
            column = machine_sheet.max_column + 1
            cell = machine_sheet.cell(1, column)
            cell.value = "Products_JSON"
            cell.fill = PatternFill("solid", fgColor="102A43")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            machine_sheet.column_dimensions[cell.column_letter].width = 64
            headers["Products_JSON"] = column
            changed = True

        for row_index in range(2, machine_sheet.max_row + 1):
            machine_id = str(machine_sheet.cell(row_index, headers["Machine_ID"]).value or "")
            existing = machine_sheet.cell(row_index, headers["Products_JSON"]).value
            if str(existing or "").strip():
                continue
            products = normalise_products(
                [],
                machine_id,
                machine_sheet.cell(row_index, headers.get("Barcodes_JSON", 0)).value if headers.get("Barcodes_JSON") else [],
                machine_sheet.cell(row_index, headers.get("Image_File", 0)).value if headers.get("Image_File") else "",
            )
            machine_sheet.cell(row_index, headers["Products_JSON"]).value = json.dumps(products, ensure_ascii=False)
            changed = True

        machine_sheet.auto_filter.ref = f"A1:{machine_sheet.cell(1, machine_sheet.max_column).coordinate}"

        settings = workbook["Settings"]
        settings_headers = {cell.value: cell.column for cell in settings[1] if cell.value}
        version_row = None
        for row_index in range(2, settings.max_row + 1):
            if str(settings.cell(row_index, settings_headers["Key"]).value or "") == "Database_Version":
                version_row = row_index
                break
        if version_row is None:
            version_row = settings.max_row + 1
            settings.cell(version_row, settings_headers["Key"]).value = "Database_Version"
            settings.cell(version_row, settings_headers["Description"]).value = "Excel database schema version"
            changed = True
        if str(settings.cell(version_row, settings_headers["Value"]).value or "") != "1.2":
            settings.cell(version_row, settings_headers["Value"]).value = "1.2"
            changed = True

        if changed:
            temp_path = path.with_suffix(path.suffix + ".v217.tmp")
            workbook.save(temp_path)
            os.replace(temp_path, path)
    finally:
        workbook.close()


def _install_database_support(namespace: dict[str, Any]) -> None:
    import database as database_module

    database_class = namespace["ExcelDatabase"]
    schema = database_module.SCHEMA
    if "Products_JSON" not in schema["Machine_Master"]:
        schema["Machine_Master"].append("Products_JSON")
    for column in ("Meter_Mode", "Manual_Coins_Used", "Play_Rule_Coins", "Win_Rate"):
        if column not in schema["Machine_Closing"]:
            schema["Machine_Closing"].append(column)
    schema["Product_Closing"] = list(PRODUCT_CLOSING_HEADERS)

    if getattr(database_class, "_v217_installed", False):
        return

    original_create = database_class.create_default_database
    original_migrate = database_class._migrate_schema
    original_save_closing = database_class.save_closing

    @staticmethod
    def create_default_database(path: Path) -> None:
        original_create(path)
        _upgrade_workbook(Path(path))

    def migrate_schema(self) -> None:
        original_migrate(self)
        _upgrade_workbook(self.path)

    def save_closing(self, header, machine_rows, user_name, product_rows=None):
        closing_id = original_save_closing(self, header, machine_rows, user_name)
        if product_rows is None:
            return closing_id

        workbook = load_workbook(self.path)
        try:
            sheet = workbook["Product_Closing"]
            headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
            delete_rows = []
            for row_index in range(2, sheet.max_row + 1):
                if str(sheet.cell(row_index, headers["Closing_ID"]).value or "") == closing_id:
                    delete_rows.append(row_index)
            for row_index in reversed(delete_rows):
                sheet.delete_rows(row_index, 1)

            now = datetime.now().isoformat(timespec="seconds")
            for product in product_rows:
                row = dict(product)
                row["Closing_ID"] = closing_id
                row["Created_At"] = now
                row["Updated_At"] = now
                row_index = sheet.max_row + 1
                for column_name in PRODUCT_CLOSING_HEADERS:
                    sheet.cell(row_index, headers[column_name]).value = row.get(column_name)
            self._save_atomic(workbook)
            return closing_id
        finally:
            workbook.close()

    def get_closing_products(self, closing_id: str) -> list[dict[str, Any]]:
        with self._workbook(read_only=True) as workbook:
            sheet = workbook["Product_Closing"]
            rows = self._sheet_rows(sheet)
        return [dict(row) for row in rows if str(row.get("Closing_ID") or "") == closing_id]

    def previous_product_states(
        self,
        machine_ids: list[str],
        before_date: date | None = None,
    ) -> dict[str, dict[str, dict[str, Any]]]:
        requested = {str(item) for item in machine_ids}
        finalized: dict[str, tuple[str, str]] = {}
        for row in self.list_closings(5000):
            if row.get("Workflow_Status") != "Finalized":
                continue
            report_date = str(row.get("Report_Date") or "")[:10]
            if before_date and report_date >= before_date.isoformat():
                continue
            finalized[str(row.get("Closing_ID") or "")] = (
                report_date,
                str(row.get("Finalized_At") or ""),
            )

        with self._workbook(read_only=True) as workbook:
            rows = self._sheet_rows(workbook["Product_Closing"])

        latest: dict[tuple[str, str], tuple[tuple[str, str], dict[str, Any]]] = {}
        for row in rows:
            machine_id = str(row.get("Machine_ID") or "")
            if machine_id not in requested:
                continue
            closing_id = str(row.get("Closing_ID") or "")
            key = finalized.get(closing_id)
            if key is None:
                continue
            product_key = str(row.get("Product_ID") or row.get("Barcode") or "").casefold()
            if not product_key:
                continue
            current = latest.get((machine_id, product_key))
            if current is None or key > current[0]:
                latest[(machine_id, product_key)] = (key, dict(row))

        result: dict[str, dict[str, dict[str, Any]]] = {}
        for (machine_id, product_key), value in latest.items():
            result.setdefault(machine_id, {})[product_key] = value[1]
            barcode_key = str(value[1].get("Barcode") or "").casefold()
            if barcode_key:
                result[machine_id][barcode_key] = value[1]
        return result

    database_class.create_default_database = create_default_database
    database_class._migrate_schema = migrate_schema
    database_class.save_closing = save_closing
    database_class.get_closing_products = get_closing_products
    database_class.previous_product_states = previous_product_states
    database_class._v217_installed = True


def _decode_image_data(image_data: Any) -> tuple[bytes, str] | None:
    value = str(image_data or "").strip()
    if not value:
        return None
    match = re.fullmatch(
        r"data:(image/(?:jpeg|png|webp));base64,([A-Za-z0-9+/=\s]+)",
        value,
        flags=re.IGNORECASE,
    )
    if not match:
        raise ValueError("Upload or paste a JPG, PNG, or WEBP product image.")
    mime_type = match.group(1).lower()
    extension = IMAGE_MIME_EXTENSIONS[mime_type]
    try:
        payload = base64.b64decode(re.sub(r"\s+", "", match.group(2)), validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("The product image is invalid.") from exc
    if not payload or len(payload) > PRODUCT_IMAGE_MAX_BYTES:
        raise ValueError("Product images must be between 1 byte and 5 MB.")
    valid = (
        mime_type == "image/jpeg" and payload.startswith(b"\xff\xd8\xff")
    ) or (
        mime_type == "image/png" and payload.startswith(b"\x89PNG\r\n\x1a\n")
    ) or (
        mime_type == "image/webp"
        and len(payload) >= 12
        and payload[:4] == b"RIFF"
        and payload[8:12] == b"WEBP"
    )
    if not valid:
        raise ValueError("The pasted or uploaded file does not match its image format.")
    return payload, extension


def _save_product_image(
    image_dir: Path,
    machine_id: str,
    product_id: str,
    image_data: Any,
    existing_file: Any,
) -> str:
    decoded = _decode_image_data(image_data)
    if decoded is None:
        return Path(str(existing_file or "")).name
    payload, extension = decoded
    safe_machine = re.sub(r"[^A-Za-z0-9_.-]+", "-", machine_id).strip("-") or "machine"
    safe_product = re.sub(r"[^A-Za-z0-9_.-]+", "-", product_id).strip("-") or "product"
    digest = hashlib.sha256(payload).hexdigest()[:12]
    file_name = f"{safe_machine}-{safe_product}-{digest}{extension}"
    image_dir.mkdir(parents=True, exist_ok=True)
    destination = image_dir / file_name
    temp = destination.with_suffix(destination.suffix + ".tmp")
    temp.write_bytes(payload)
    os.replace(temp, destination)
    return file_name


def _remove_product_image(image_dir: Path, file_name: Any, legacy_machine_image: Any = "") -> None:
    safe_name = Path(str(file_name or "")).name
    legacy_name = Path(str(legacy_machine_image or "")).name
    if not safe_name or safe_name == legacy_name:
        return
    try:
        (image_dir / safe_name).unlink(missing_ok=True)
    except OSError:
        pass


def _remove_routes(api, replacements: list[tuple[str, set[str]]]) -> None:
    retained = []
    for route in api.router.routes:
        path = getattr(route, "path", "")
        methods = set(getattr(route, "methods", set()) or set())
        if any(path == target and methods.intersection(target_methods) for target, target_methods in replacements):
            continue
        retained.append(route)
    api.router.routes[:] = retained


def _install_api_support(namespace: dict[str, Any]) -> None:
    api = namespace["api"]
    db: Callable[[], Any] = namespace["db"]
    serialise = namespace["serialise"]
    parse_iso_date = namespace["parse_iso_date"]
    make_sales = namespace["make_sales"]
    calculate_closing = namespace["calculate_closing"]
    MachineInput = namespace["MachineInput"]
    APP_NAME = namespace["APP_NAME"]
    APP_VERSION = namespace["APP_VERSION"]
    DEFAULT_OUTLET = namespace["DEFAULT_OUTLET"]
    REPORTS_DIR: Path = namespace["REPORTS_DIR"]
    MACHINE_IMAGE_DIR: Path = namespace["MACHINE_IMAGE_DIR"]
    export_daily_closing = namespace["export_daily_closing"]
    application_settings = namespace["application_settings"]
    canonical_machine_type_name = namespace["canonical_machine_type_name"]
    machine_type_rule = namespace["machine_type_rule"]

    replacements = [
        ("/api/bootstrap", {"GET"}),
        ("/api/machines", {"GET", "POST"}),
        ("/api/machines/{machine_id}", {"DELETE"}),
        ("/api/new-closing", {"GET"}),
        ("/api/closings/{closing_id}", {"GET"}),
        ("/api/calculate", {"POST"}),
        ("/api/closings/save", {"POST"}),
        ("/api/reports/pdf", {"POST"}),
    ]
    _remove_routes(api, replacements)

    def machine_input_bundle(payload_rows: list[dict[str, Any]], finalized: bool = False):
        inputs = []
        product_records: list[dict[str, Any]] = []
        metadata: list[dict[str, Any]] = []
        for row in payload_rows:
            machine_id = str(row.get("machine_id") or "").strip()
            machine_name = str(row.get("machine_name") or "").strip()
            products = list(row.get("products") or [])
            if finalized and not products:
                raise ValueError(f"{machine_name}: add at least one product in Daily Closing → Edit.")
            begin_total = final_total = 0
            for index, product in enumerate(products):
                product_id = _safe_product_id(product.get("product_id"), machine_id, str(product.get("barcode") or ""), index)
                barcode = str(product.get("barcode") or "").strip()
                if finalized and not barcode:
                    raise ValueError(f"{machine_name}: every product requires a barcode before finalizing.")
                begin_qty = _to_int(product.get("begin_qty"))
                final_raw = product.get("final_qty")
                final_present = str(final_raw if final_raw is not None else "").strip() != ""
                if finalized and not final_present:
                    raise ValueError(f"{machine_name} · {barcode or 'product'}: Final Qty is required before finalizing.")
                final_qty = _to_int(final_raw) if final_present else begin_qty
                if begin_qty < 0 or (final_present and final_qty < 0):
                    raise ValueError(f"{machine_name}: product quantities cannot be negative.")
                available = begin_qty
                if final_present and final_qty > available:
                    raise ValueError(f"{machine_name} · {barcode or 'product'}: Final Qty cannot exceed available quantity.")
                qty_used = available - final_qty if final_present else 0
                begin_total += begin_qty
                final_total += final_qty
                product_records.append(
                    {
                        "Report_Date": "",
                        "Machine_ID": machine_id,
                        "Machine_Name": machine_name,
                        "Product_ID": product_id,
                        "Barcode": barcode,
                        "Image_File": Path(str(product.get("image_file") or "")).name,
                        "Begin_Qty": begin_qty,
                        "Final_Qty": final_qty if final_present else None,
                        "Qty_Used": qty_used if final_present else None,
                    }
                )

            begin_raw = row.get("begin_coin_meter")
            final_raw = row.get("final_coin_meter")
            manual_raw = row.get("coins_used")
            begin_present = str(begin_raw if begin_raw is not None else "").strip() != ""
            final_present = str(final_raw if final_raw is not None else "").strip() != ""
            if begin_present != final_present:
                raise ValueError(f"{machine_name}: enter both Begin Meter and Final Meter, or leave both blank for manual Coins Used.")
            if begin_present:
                begin_meter = _to_int(begin_raw)
                final_meter = _to_int(final_raw)
                if final_meter < begin_meter:
                    raise ValueError(f"{machine_name}: Final Meter cannot be lower than Begin Meter.")
                coins_used = final_meter - begin_meter
                meter_mode = "meter"
                synthetic_begin = begin_meter
                synthetic_final = final_meter
            else:
                manual_present = str(manual_raw if manual_raw is not None else "").strip() != ""
                if finalized and not manual_present:
                    raise ValueError(f"{machine_name}: enter Coins Used manually or enter both meter values.")
                coins_used = _to_int(manual_raw)
                if coins_used < 0:
                    raise ValueError(f"{machine_name}: Coins Used cannot be negative.")
                meter_mode = "manual"
                synthetic_begin = 0
                synthetic_final = coins_used

            inputs.append(
                MachineInput(
                    machine_id=machine_id,
                    machine_name=machine_name,
                    machine_type=str(row.get("machine_type") or "").strip(),
                    capacity=max(1, _to_int(row.get("capacity")), begin_total),
                    begin_prize=begin_total,
                    refill_prize=0,
                    final_prize=final_total,
                    begin_coin_meter=synthetic_begin,
                    final_coin_meter=synthetic_final,
                    status=str(row.get("status") or "Working"),
                    notes=str(row.get("notes") or "").strip(),
                )
            )
            metadata.append(
                {
                    "machine_id": machine_id,
                    "begin_meter": _to_int(begin_raw) if begin_present else None,
                    "final_meter": _to_int(final_raw) if final_present else None,
                    "coins_used": coins_used,
                    "meter_mode": meter_mode,
                }
            )
        return inputs, product_records, metadata

    def machine_view_v213(record: dict[str, Any]) -> dict[str, Any]:
        view = machine_product_view(record)
        view["Machine_Type"] = canonical_machine_type_name(view.get("Machine_Type"), str(view.get("Machine_ID") or ""))
        return view

    def bootstrap_v217() -> dict[str, Any]:
        database = db()
        return serialise(
            {
                "app": {"name": APP_NAME, "version": APP_VERSION},
                "settings": {
                    "outlet": database.get_setting("Outlet", DEFAULT_OUTLET),
                    "currency": database.get_setting("Currency", "USD"),
                    "variance_tolerance": _to_int(database.get_setting("Variance_Tolerance", 0)),
                    **application_settings(database),
                },
                "dashboard": database.dashboard_summary(),
                "recent_closings": database.list_closings(12),
                "machines": [machine_view_v213(item) for item in database.list_machines(active_only=False)],
            }
        )

    def list_machines_v217(active_only: bool = False) -> list[dict[str, Any]]:
        return serialise([machine_view_v213(item) for item in db().list_machines(active_only=active_only)])

    def save_machine_v217(payload: dict[str, Any]) -> dict[str, Any]:
        database = db()
        all_machines = database.list_machines(active_only=False)
        machine_id = str(payload.get("machine_id") or "").strip().upper()
        if machine_id:
            existing = next(
                (item for item in all_machines if str(item.get("Machine_ID") or "").upper() == machine_id),
                {},
            )
        else:
            machine_id = f"MCH-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
            while any(str(item.get("Machine_ID") or "").upper() == machine_id for item in all_machines):
                machine_id = f"MCH-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
            existing = {}
        existing_products = normalise_products(
            existing.get("Products_JSON"),
            machine_id,
            existing.get("Barcodes_JSON"),
            existing.get("Image_File"),
        )
        existing_by_id = {str(item.get("product_id") or "").casefold(): item for item in existing_products}
        existing_by_code = {str(item.get("barcode") or "").casefold(): item for item in existing_products if item.get("barcode")}

        incoming = payload.get("products")
        if not isinstance(incoming, list):
            incoming = [{"barcode": code} for code in _parse_json_list(payload.get("barcodes"))]
        saved_products: list[dict[str, Any]] = []
        seen_codes: set[str] = set()
        seen_ids: set[str] = set()
        legacy_image = existing.get("Image_File")
        for index, source in enumerate(incoming):
            if not isinstance(source, dict):
                source = {"barcode": source}
            barcode = str(source.get("barcode") or "").strip()
            default_quantity = 0
            requested_id = _safe_product_id(source.get("product_id") or source.get("id"), machine_id, barcode, index)
            matched = existing_by_id.get(requested_id.casefold()) or existing_by_code.get(barcode.casefold()) or {}
            product_id = str(matched.get("product_id") or requested_id)
            if product_id.casefold() in seen_ids:
                product_id = _safe_product_id("", machine_id, f"{barcode}-{index}", index)
            seen_ids.add(product_id.casefold())
            if barcode:
                key = barcode.casefold()
                if key in seen_codes:
                    raise ValueError(f"Duplicate barcode: {barcode}")
                seen_codes.add(key)
            existing_image = matched.get("image_file") or ""
            has_new_image = bool(str(source.get("image_data") or "").strip())
            remove_requested = bool(source.get("remove_image"))
            if not barcode and (default_quantity > 0 or has_new_image or (existing_image and not remove_requested)):
                raise ValueError(f"Product {index + 1}: Barcode is required when a quantity or image is provided.")
            if remove_requested:
                _remove_product_image(MACHINE_IMAGE_DIR, existing_image, legacy_image)
                image_file = ""
            else:
                image_file = _save_product_image(
                    MACHINE_IMAGE_DIR,
                    machine_id,
                    product_id,
                    source.get("image_data"),
                    existing_image,
                )
            if not barcode and not image_file and default_quantity == 0:
                continue
            saved_products.append(
                {
                    "product_id": product_id,
                    "barcode": barcode,
                    "default_quantity": default_quantity,
                    "image_file": image_file,
                }
            )
        if len(saved_products) > 30:
            raise ValueError("A machine can contain up to 30 products.")

        removed_ids = {str(item.get("product_id") or "").casefold() for item in existing_products} - {
            str(item.get("product_id") or "").casefold() for item in saved_products
        }
        for item in existing_products:
            if str(item.get("product_id") or "").casefold() in removed_ids:
                _remove_product_image(MACHINE_IMAGE_DIR, item.get("image_file"), legacy_image)

        existing_type = canonical_machine_type_name(existing.get("Machine_Type"), machine_id)
        requested_type = canonical_machine_type_name(payload.get("machine_type") or existing_type, machine_id)
        if not requested_type:
            raise ValueError("Machine Type is required.")
        configured_types = application_settings(database).get("machine_types", [])
        configured_match = next(
            (str(item.get("name") or "").strip() for item in configured_types if str(item.get("name") or "").strip().casefold() == requested_type.casefold()),
            "",
        )
        if configured_match:
            requested_type = configured_match
        elif not (existing and requested_type.casefold() == existing_type.casefold()):
            raise ValueError("Choose a Machine Type configured in Settings.")
        requested_name = str(payload.get("machine_name") or "").strip()
        machine_name = str(existing.get("Machine_Name") or "").strip() if existing else ""
        machine_name = machine_name or requested_name or requested_type
        same_type_rows = [
            item for item in all_machines
            if str(item.get("Machine_ID") or "").upper() != machine_id
            and canonical_machine_type_name(item.get("Machine_Type"), str(item.get("Machine_ID") or "")).casefold() == requested_type.casefold()
        ]
        max_sort_order = max((_to_int(item.get("Sort_Order"), 0) for item in same_type_rows), default=0)
        if existing and requested_type.casefold() == existing_type.casefold():
            auto_sort_order = _to_int(existing.get("Sort_Order"), 0) or (max_sort_order + 1)
        else:
            auto_sort_order = max_sort_order + 1
        record = {
            "Machine_ID": machine_id,
            "Machine_Name": machine_name,
            "Machine_Type": requested_type,
            "Capacity": 0,
            "Prize_Category": str(payload.get("prize_category") or ""),
            "Active": bool(payload.get("active", True)),
            "Sort_Order": auto_sort_order,
            "Notes": str(payload.get("notes") or "").strip(),
            "Barcodes_JSON": json.dumps([item["barcode"] for item in saved_products if item.get("barcode")], ensure_ascii=False),
            "Image_File": str(existing.get("Image_File") or ""),
            "Products_JSON": json.dumps(saved_products, ensure_ascii=False),
        }
        database.upsert_machine(record, str(payload.get("user_name") or "Administrator"))
        return {"ok": True, "machine": serialise(machine_view_v213(record))}

    def delete_machine_v215(machine_id: str) -> dict[str, Any]:
        database = db()
        target_id = str(machine_id or "").strip().upper()
        existing = next(
            (item for item in database.list_machines(active_only=False) if str(item.get("Machine_ID") or "").upper() == target_id),
            None,
        )
        if not existing:
            raise HTTPException(status_code=404, detail="Machine was not found.")
        record = dict(existing)
        record["Active"] = False
        database.upsert_machine(record, "Administrator")
        return {"ok": True, "machine": serialise(machine_view_v213(record)), "deactivated": True}

    def new_closing_v217(report_date: str) -> dict[str, Any]:
        requested_date = parse_iso_date(report_date)
        database = db()
        machines = database.list_machines(active_only=True)
        machine_ids = [str(item.get("Machine_ID") or "") for item in machines]
        # Carry forward the latest finalized state up to and including the selected report date.
        # The state helpers use an exclusive cutoff, so use the following day as the cutoff.
        carry_cutoff = requested_date.fromordinal(requested_date.toordinal() + 1)
        machine_states = database.previous_machine_states(machine_ids, carry_cutoff)
        product_states = database.previous_product_states(machine_ids, carry_cutoff)
        rows = []
        for machine in machines:
            machine_id = str(machine.get("Machine_ID") or "")
            previous_machine = machine_states.get(machine_id, {})
            previous_products = product_states.get(machine_id, {})
            products = normalise_products(
                machine.get("Products_JSON"),
                machine_id,
                machine.get("Barcodes_JSON"),
                machine.get("Image_File"),
            )
            product_rows = []
            for product in products:
                previous = previous_products.get(str(product.get("product_id") or "").casefold()) or previous_products.get(str(product.get("barcode") or "").casefold()) or {}
                begin_qty = _to_int(previous.get("Final_Qty"), 0)
                product_rows.append(
                    {
                        **product_view(product),
                        "begin_qty": begin_qty,
                        "final_qty": "",
                    }
                )
            meter_mode = str(previous_machine.get("Meter_Mode") or "meter")
            begin_meter: Any = ""
            final_meter: Any = ""
            if previous_machine and meter_mode != "manual" and previous_machine.get("Final_Coin_Meter") is not None:
                begin_meter = _to_int(previous_machine.get("Final_Coin_Meter"))
                final_meter = begin_meter
            rows.append(
                {
                    "machine_id": machine_id,
                    "machine_name": machine.get("Machine_Name"),
                    "machine_type": canonical_machine_type_name(machine.get("Machine_Type"), machine_id),
                    "sort_order": _to_int(machine.get("Sort_Order"), 0),
                    "capacity": 0,
                    "products": product_rows,
                    "barcodes": [item.get("barcode") for item in product_rows if item.get("barcode")],
                    "begin_coin_meter": begin_meter,
                    "final_coin_meter": final_meter,
                    "coins_used": "",
                    "status": "Working",
                    "notes": "",
                }
            )
        return serialise(
            {
                "report_date": requested_date.isoformat(),
                "outlet": database.get_setting("Outlet", DEFAULT_OUTLET),
                "machines": rows,
            }
        )

    def get_closing_v217(closing_id: str) -> dict[str, Any]:
        try:
            header, machines = db().get_closing(closing_id)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        products = db().get_closing_products(closing_id)
        return serialise({"header": header, "machines": machines, "products": products})

    def calculate_v217(payload: dict[str, Any]) -> dict[str, Any]:
        sales = make_sales(dict(payload.get("sales") or {}))
        machines, product_records, metadata = machine_input_bundle(list(payload.get("machines") or []), False)
        tolerance = _to_int(db().get_setting("Variance_Tolerance", 0))
        result = calculate_closing(sales, machines, tolerance)
        return serialise(
            {
                "summary": result.as_record(),
                "machines": [item.as_record() for item in result.machines],
                "products": product_records,
                "meter_modes": metadata,
            }
        )

    def save_closing_v217(payload: dict[str, Any]) -> dict[str, Any]:
        database = db()
        report_date_value = parse_iso_date(payload.get("report_date"))
        workflow_status = str(payload.get("workflow_status") or "Draft").strip().title()
        if workflow_status not in {"Draft", "Finalized"}:
            raise ValueError("Workflow status must be Draft or Finalized.")
        closed_by = str(payload.get("closed_by") or "").strip()
        verified_by = str(payload.get("verified_by") or "").strip()
        if workflow_status == "Finalized" and not closed_by:
            raise ValueError("Closed By is required before finalizing.")
        if workflow_status == "Finalized" and not verified_by:
            raise ValueError("Verified By is required before finalizing.")

        sales_payload = dict(payload.get("sales") or {})
        sales = make_sales(sales_payload)
        payload_machines = list(payload.get("machines") or [])
        machines, product_records, metadata = machine_input_bundle(payload_machines, workflow_status == "Finalized")
        tolerance = _to_int(database.get_setting("Variance_Tolerance", 0))
        result = calculate_closing(sales, machines, tolerance)
        if workflow_status == "Finalized" and result.coin_variance != 0:
            raise ValueError(
                "The closing cannot be finalized until the coin variance is zero. "
                f"Current variance: {result.coin_variance:+,} coins."
            )

        closing_id = str(payload.get("closing_id") or "").strip() or database.next_closing_id(report_date_value)
        settings_now = application_settings(database)
        try:
            exchange_rate = float(sales_payload.get("exchange_rate_usd_khr") or settings_now.get("exchange_rate_usd_khr") or 4100)
            coin_price = float(sales_payload.get("price_per_coin_usd") or settings_now.get("price_per_coin_usd") or 0.3125)
            cash_sales_khr = float(sales_payload.get("cash_sales_khr") or (float(sales.cash_sales) * exchange_rate))
        except (TypeError, ValueError) as exc:
            raise ValueError("Exchange rate, price per coin, and cash sales must be valid numbers.") from exc
        if exchange_rate <= 0 or coin_price <= 0:
            raise ValueError("Exchange rate and price per coin must be greater than 0.")
        coins_used = int(result.coins_dispensed)
        coin_return = int(result.machine_coins_used)
        lose_over = int(result.coin_variance)
        expected_coin_revenue = float(coins_used) * coin_price
        discount_usd = round(expected_coin_revenue - float(result.total_sales), 2)
        discount_percent = (discount_usd / expected_coin_revenue * 100.0) if expected_coin_revenue > 0 else None
        total_plays = 0.0
        for item in result.machines:
            rule = machine_type_rule(database, item.machine_type, item.machine_id)
            if rule and int(rule[1]) > 0 and int(item.coins_used) > 0:
                total_plays += float(item.coins_used) / float(rule[1])
        overall_win_rate = (total_plays / float(result.total_prizes_won)) if result.total_prizes_won > 0 and total_plays > 0 else None
        header = {
            "Closing_ID": closing_id,
            "Report_Date": report_date_value.isoformat(),
            "Outlet": str(payload.get("outlet") or database.get_setting("Outlet", DEFAULT_OUTLET)).strip(),
            "Cash_Sales": float(sales.cash_sales),
            "Cash_Transactions": sales.cash_transactions,
            "ABA_Sales": float(sales.aba_sales),
            "ABA_Transactions": sales.aba_transactions,
            "Adjustment": float(sales.adjustment),
            "Beginning_Coins": sales.beginning_coins,
            "Coins_Added": sales.coins_added,
            "Final_Coins": sales.final_coins,
            "Coins_Dispensed": result.coins_dispensed,
            "Total_Sales": float(result.total_sales),
            "Total_Transactions": result.total_transactions,
            "Average_Sale_Value_Per_Coin": float(result.average_sale_value_per_coin),
            "Machine_Coins_Used": result.machine_coins_used,
            "Coin_Variance": result.coin_variance,
            "Total_Prizes_Won": result.total_prizes_won,
            "Average_Coins_Per_Prize": float(result.average_coins_per_prize),
            "Average_Revenue_Per_Prize": float(result.average_revenue_per_prize),
            "Cash_Sales_KHR": cash_sales_khr,
            "Exchange_Rate_USD_KHR": exchange_rate,
            "Price_Per_Coin_USD": coin_price,
            "Coins_Used": coins_used,
            "Coin_Return": coin_return,
            "Lose_Over": lose_over,
            "Discount_USD": discount_usd,
            "Discount_Percent": discount_percent,
            "Overall_Win_Rate": overall_win_rate,
            "Closing_Status": result.closing_status,
            "Workflow_Status": workflow_status,
            "Closed_By": closed_by,
            "Verified_By": verified_by,
            "Notes": str(payload.get("notes") or "").strip(),
        }
        metadata_by_id = {item["machine_id"]: item for item in metadata}
        machine_records = []
        for item in result.machines:
            meta = metadata_by_id[item.machine_id]
            rule = machine_type_rule(database, item.machine_type, item.machine_id)
            play_rule_coins = int(rule[1]) if rule else 0
            win_rate = None
            if play_rule_coins > 0 and item.prizes_won > 0:
                win_rate = float(meta["coins_used"]) / float(play_rule_coins) / float(item.prizes_won)
            machine_records.append(
                {
                    "Machine_ID": item.machine_id,
                    "Machine_Name": item.machine_name,
                    "Machine_Type": canonical_machine_type_name(item.machine_type, item.machine_id),
                    "Capacity": item.capacity,
                    "Begin_Prize": item.begin_prize,
                    "Refill_Prize": item.refill_prize,
                    "Available_Prize": item.available_prize,
                    "Final_Prize": item.final_prize,
                    "Prizes_Won": item.prizes_won,
                    "Refill_Needed": item.refill_needed,
                    "Begin_Coin_Meter": meta["begin_meter"],
                    "Final_Coin_Meter": meta["final_meter"],
                    "Coins_Used": meta["coins_used"],
                    "Coins_Per_Prize": float(item.coins_per_prize),
                    "Play_Rule_Coins": play_rule_coins or None,
                    "Win_Rate": win_rate,
                    "Allocated_Sales": float(item.allocated_sales),
                    "Average_Revenue_Per_Prize": float(item.avg_revenue_per_prize),
                    "Machine_Status": item.status,
                    "Notes": item.notes,
                    "Meter_Mode": meta["meter_mode"],
                    "Manual_Coins_Used": meta["coins_used"] if meta["meter_mode"] == "manual" else None,
                }
            )
        for product in product_records:
            product["Report_Date"] = report_date_value.isoformat()
        database.save_closing(header, machine_records, closed_by, product_rows=product_records)

        report_path = None
        if workflow_status == "Finalized":
            saved_header, saved_machines = database.get_closing(closing_id)
            REPORTS_DIR.mkdir(parents=True, exist_ok=True)
            report_path = REPORTS_DIR / f"{closing_id}.xlsx"
            export_daily_closing(saved_header, saved_machines, report_path)
            _append_product_detail_sheet(report_path, product_records)

        return serialise(
            {
                "ok": True,
                "closing_id": closing_id,
                "workflow_status": workflow_status,
                "result": result.as_record(),
                "report_path": str(report_path) if report_path else None,
            }
        )

    def export_current_pdf(payload: dict[str, Any]) -> dict[str, Any]:
        database = db()
        report_date_value = parse_iso_date(payload.get("report_date"))
        sales_payload = dict(payload.get("sales") or {})
        sales = make_sales(sales_payload)
        payload_machines = list(payload.get("machines") or [])
        machines, product_records, metadata = machine_input_bundle(payload_machines, False)
        tolerance = _to_int(database.get_setting("Variance_Tolerance", 0))
        result = calculate_closing(sales, machines, tolerance)

        settings_now = application_settings(database)
        try:
            exchange_rate = float(sales_payload.get("exchange_rate_usd_khr") or settings_now.get("exchange_rate_usd_khr") or 4100)
            coin_price = float(sales_payload.get("price_per_coin_usd") or settings_now.get("price_per_coin_usd") or 0.3125)
            cash_sales_khr = float(sales_payload.get("cash_sales_khr") or (float(sales.cash_sales) * exchange_rate))
        except (TypeError, ValueError) as exc:
            raise ValueError("Exchange rate, price per coin, and cash sales must be valid numbers.") from exc
        if exchange_rate <= 0 or coin_price <= 0:
            raise ValueError("Exchange rate and price per coin must be greater than 0.")

        coins_used = int(result.coins_dispensed)
        coin_return = int(result.machine_coins_used)
        lose_over = int(result.coin_variance)
        expected_coin_revenue = float(coins_used) * coin_price
        discount_usd = round(expected_coin_revenue - float(result.total_sales), 2)
        discount_percent = (discount_usd / expected_coin_revenue * 100.0) if expected_coin_revenue > 0 else None

        metadata_by_id = {item["machine_id"]: item for item in metadata}
        products_by_machine: dict[str, list[dict[str, Any]]] = {}
        for product in product_records:
            products_by_machine.setdefault(str(product.get("Machine_ID") or ""), []).append(
                {
                    "barcode": product.get("Barcode") or "",
                    "begin_qty": product.get("Begin_Qty") or 0,
                    "final_qty": product.get("Final_Qty") or 0,
                    "qty_used": product.get("Qty_Used") or 0,
                }
            )

        total_plays = 0.0
        machine_rows: list[dict[str, Any]] = []
        for item in result.machines:
            meta = metadata_by_id[item.machine_id]
            rule = machine_type_rule(database, item.machine_type, item.machine_id)
            play_rule_coins = int(rule[1]) if rule else 0
            machine_win_rate = None
            if play_rule_coins > 0 and item.prizes_won > 0:
                machine_win_rate = float(meta["coins_used"]) / float(play_rule_coins) / float(item.prizes_won)
            if play_rule_coins > 0 and int(meta["coins_used"]) > 0:
                total_plays += float(meta["coins_used"]) / float(play_rule_coins)
            machine_rows.append(
                {
                    "machine_id": item.machine_id,
                    "machine_type": canonical_machine_type_name(item.machine_type, item.machine_id),
                    "products": products_by_machine.get(item.machine_id, []),
                    "begin_meter": meta.get("begin_meter"),
                    "final_meter": meta.get("final_meter"),
                    "coins_used": meta.get("coins_used") or 0,
                    "win_rate": machine_win_rate,
                    "status": item.status,
                }
            )

        overall_win_rate = (total_plays / float(result.total_prizes_won)) if result.total_prizes_won > 0 and total_plays > 0 else None
        avg_product = float(result.total_sales) / float(result.total_prizes_won) if result.total_prizes_won > 0 else 0.0
        report = {
            "report_date": report_date_value.isoformat(),
            "outlet": str(payload.get("outlet") or database.get_setting("Outlet", DEFAULT_OUTLET)).strip(),
            "closed_by": str(payload.get("closed_by") or "").strip(),
            "verified_by": str(payload.get("verified_by") or "").strip(),
            "notes": str(payload.get("notes") or "").strip(),
            "cash_sales_khr": cash_sales_khr,
            "aba_sales": float(sales.aba_sales),
            "exchange_rate": exchange_rate,
            "coin_price": coin_price,
            "beginning_coins": sales.beginning_coins,
            "coins_added": sales.coins_added,
            "final_coins": sales.final_coins,
            "total_sales": float(result.total_sales),
            "coins_used": coins_used,
            "coin_return": coin_return,
            "lose_over": lose_over,
            "products_used": result.total_prizes_won,
            "avg_product": avg_product,
            "discount_usd": discount_usd,
            "discount_percent": discount_percent,
            "overall_win_rate": overall_win_rate,
            "machines": machine_rows,
        }

        closing_id = str(payload.get("closing_id") or "").strip()
        if closing_id:
            safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", closing_id)
        else:
            safe_name = f"Daily_Closing_{report_date_value.strftime('%Y%m%d')}"
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        destination = REPORTS_DIR / f"{safe_name}.pdf"
        build_daily_closing_pdf(report, destination)
        try:
            if os.name == "nt":
                os.startfile(destination)  # type: ignore[attr-defined]
            else:
                import subprocess
                import sys
                subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(destination)])
        except Exception:
            # The PDF is still successfully created even if the OS cannot open it.
            pass
        return {"ok": True, "path": str(destination), "filename": destination.name}

    api.add_api_route("/api/bootstrap", bootstrap_v217, methods=["GET"])
    api.add_api_route("/api/machines", list_machines_v217, methods=["GET"])
    api.add_api_route("/api/machines", save_machine_v217, methods=["POST"])
    api.add_api_route("/api/machines/{machine_id}", delete_machine_v215, methods=["DELETE"])
    api.add_api_route("/api/new-closing", new_closing_v217, methods=["GET"])
    api.add_api_route("/api/closings/{closing_id}", get_closing_v217, methods=["GET"])
    api.add_api_route("/api/calculate", calculate_v217, methods=["POST"])
    api.add_api_route("/api/closings/save", save_closing_v217, methods=["POST"])
    api.add_api_route("/api/reports/pdf", export_current_pdf, methods=["POST"])


def _append_product_detail_sheet(path: Path, products: list[dict[str, Any]]) -> None:
    if not path.exists():
        return
    workbook = load_workbook(path)
    try:
        if "Product Detail" in workbook.sheetnames:
            del workbook["Product Detail"]
        sheet = workbook.create_sheet("Product Detail")
        headers = ["Machine ID", "Machine Name", "Barcode", "Begin Qty", "Final Qty", "Qty Used"]
        sheet.append(headers)
        _style_header(sheet)
        for product in products:
            sheet.append(
                [
                    product.get("Machine_ID"),
                    product.get("Machine_Name"),
                    product.get("Barcode"),
                    product.get("Begin_Qty"),
                    product.get("Final_Qty"),
                    product.get("Qty_Used"),
                ]
            )
        widths = [16, 26, 24, 14, 14, 14]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
        workbook.save(path)
    finally:
        workbook.close()


def install_v217(namespace: dict[str, Any]) -> None:
    _install_database_support(namespace)
    _install_api_support(namespace)


# v2.1.39 - audited refill persistence + startup migration hotfix.
def _normalise_refill_history(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        raw = value
    else:
        text = str(value or "").strip()
        if not text:
            return []
        try:
            raw = json.loads(text)
        except (json.JSONDecodeError, TypeError):
            return []
    if not isinstance(raw, list):
        return []
    result: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        qty = max(0, _to_int(item.get("qty")))
        if qty <= 0:
            continue
        result.append(
            {
                "qty": qty,
                "at": str(item.get("at") or item.get("timestamp") or "")[:40],
                "by": str(item.get("by") or item.get("staff") or "")[:100],
            }
        )
    return result[-100:]


def _refill_total(product: dict[str, Any]) -> int:
    history = _normalise_refill_history(product.get("refill_history") or product.get("Refill_History_JSON"))
    if history:
        return sum(max(0, _to_int(item.get("qty"))) for item in history)
    return max(0, _to_int(product.get("refill_qty") or product.get("Refill_Qty")))


def _upgrade_refill_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    changed = False
    try:
        if "Product_Closing" not in workbook.sheetnames:
            return
        sheet = workbook["Product_Closing"]
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
        for column_name, width in (("Refill_Qty", 14), ("Refill_History_JSON", 60)):
            if column_name not in headers:
                column = sheet.max_column + 1
                cell = sheet.cell(1, column)
                cell.value = column_name
                cell.fill = PatternFill("solid", fgColor="102A43")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                sheet.column_dimensions[cell.column_letter].width = width
                headers[column_name] = column
                changed = True
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, sheet.max_column).coordinate}"

        if "Settings" in workbook.sheetnames:
            settings = workbook["Settings"]
            settings_headers = {cell.value: cell.column for cell in settings[1] if cell.value}
            key_col = settings_headers.get("Key")
            value_col = settings_headers.get("Value")
            desc_col = settings_headers.get("Description")
            if key_col and value_col:
                version_row = None
                for row_index in range(2, settings.max_row + 1):
                    if str(settings.cell(row_index, key_col).value or "") == "Database_Version":
                        version_row = row_index
                        break
                if version_row is None:
                    version_row = settings.max_row + 1
                    settings.cell(version_row, key_col).value = "Database_Version"
                    if desc_col:
                        settings.cell(version_row, desc_col).value = "Excel database schema version"
                if str(settings.cell(version_row, value_col).value or "") != "1.3":
                    settings.cell(version_row, value_col).value = "1.3"
                    changed = True
        if changed:
            temp_path = path.with_suffix(path.suffix + ".v2138.tmp")
            workbook.save(temp_path)
            os.replace(temp_path, path)
    finally:
        workbook.close()


def _refill_records(payload_machines: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for machine in payload_machines:
        machine_id = str(machine.get("machine_id") or "").strip()
        machine_name = str(machine.get("machine_name") or "").strip()
        for index, product in enumerate(list(machine.get("products") or [])):
            barcode = str(product.get("barcode") or "").strip()
            product_id = _safe_product_id(product.get("product_id"), machine_id, barcode, index)
            history = _normalise_refill_history(product.get("refill_history") or product.get("Refill_History_JSON"))
            refill_qty = sum(item["qty"] for item in history) if history else max(0, _to_int(product.get("refill_qty") or product.get("Refill_Qty")))
            begin_qty = max(0, _to_int(product.get("begin_qty")))
            final_raw = product.get("final_qty")
            final_present = str(final_raw if final_raw is not None else "").strip() != ""
            final_qty = max(0, _to_int(final_raw)) if final_present else begin_qty + refill_qty
            records.append(
                {
                    "machine_id": machine_id,
                    "machine_name": machine_name,
                    "product_id": product_id,
                    "barcode": barcode,
                    "begin_qty": begin_qty,
                    "refill_qty": refill_qty,
                    "final_qty": final_qty,
                    "final_present": final_present,
                    "refill_history": history,
                }
            )
    return records


def _payload_with_refill_available(payload: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    copied = dict(payload)
    copied_machines: list[dict[str, Any]] = []
    originals = _refill_records(list(payload.get("machines") or []))
    original_by_key = {(item["machine_id"], item["product_id"].casefold()): item for item in originals}
    original_by_barcode = {(item["machine_id"], item["barcode"].casefold()): item for item in originals if item["barcode"]}
    for machine in list(payload.get("machines") or []):
        machine_copy = dict(machine)
        products_copy: list[dict[str, Any]] = []
        machine_id = str(machine.get("machine_id") or "").strip()
        for index, product in enumerate(list(machine.get("products") or [])):
            product_copy = dict(product)
            barcode = str(product.get("barcode") or "").strip()
            product_id = _safe_product_id(product.get("product_id"), machine_id, barcode, index)
            original = original_by_key.get((machine_id, product_id.casefold())) or original_by_barcode.get((machine_id, barcode.casefold()))
            refill_qty = int(original["refill_qty"]) if original else _refill_total(product_copy)
            product_copy["begin_qty"] = max(0, _to_int(product.get("begin_qty"))) + refill_qty
            products_copy.append(product_copy)
        machine_copy["products"] = products_copy
        copied_machines.append(machine_copy)
    copied["machines"] = copied_machines
    return copied, originals


def _patch_saved_refills(database, closing_id: str, records: list[dict[str, Any]]) -> None:
    if not closing_id:
        return
    workbook = load_workbook(database.path)
    try:
        product_sheet = workbook["Product_Closing"]
        product_headers = {cell.value: cell.column for cell in product_sheet[1] if cell.value}
        record_by_id = {(item["machine_id"], item["product_id"].casefold()): item for item in records}
        record_by_barcode = {(item["machine_id"], item["barcode"].casefold()): item for item in records if item["barcode"]}
        for row_index in range(2, product_sheet.max_row + 1):
            if str(product_sheet.cell(row_index, product_headers["Closing_ID"]).value or "") != closing_id:
                continue
            machine_id = str(product_sheet.cell(row_index, product_headers["Machine_ID"]).value or "")
            product_id = str(product_sheet.cell(row_index, product_headers["Product_ID"]).value or "").casefold()
            barcode = str(product_sheet.cell(row_index, product_headers["Barcode"]).value or "").casefold()
            record = record_by_id.get((machine_id, product_id)) or record_by_barcode.get((machine_id, barcode))
            if not record:
                continue
            available = int(record["begin_qty"]) + int(record["refill_qty"])
            final_present = bool(record.get("final_present", True))
            final_qty = int(record["final_qty"]) if final_present else available
            product_sheet.cell(row_index, product_headers["Begin_Qty"]).value = int(record["begin_qty"])
            product_sheet.cell(row_index, product_headers["Refill_Qty"]).value = int(record["refill_qty"])
            if not final_present:
                product_sheet.cell(row_index, product_headers["Final_Qty"]).value = None
            product_sheet.cell(row_index, product_headers["Qty_Used"]).value = (available - final_qty) if final_present else None
            product_sheet.cell(row_index, product_headers["Refill_History_JSON"]).value = json.dumps(record["refill_history"], ensure_ascii=False)

        if "Machine_Closing" in workbook.sheetnames:
            machine_sheet = workbook["Machine_Closing"]
            machine_headers = {cell.value: cell.column for cell in machine_sheet[1] if cell.value}
            machine_totals: dict[str, dict[str, int]] = {}
            for item in records:
                totals = machine_totals.setdefault(item["machine_id"], {"begin": 0, "refill": 0, "final": 0})
                totals["begin"] += int(item["begin_qty"])
                totals["refill"] += int(item["refill_qty"])
                totals["final"] += int(item["final_qty"])
            closing_col = machine_headers.get("Closing_ID")
            machine_col = machine_headers.get("Machine_ID")
            if closing_col and machine_col:
                for row_index in range(2, machine_sheet.max_row + 1):
                    if str(machine_sheet.cell(row_index, closing_col).value or "") != closing_id:
                        continue
                    machine_id = str(machine_sheet.cell(row_index, machine_col).value or "")
                    totals = machine_totals.get(machine_id)
                    if not totals:
                        continue
                    for name, value in (
                        ("Begin_Prize", totals["begin"]),
                        ("Refill_Prize", totals["refill"]),
                        ("Available_Prize", totals["begin"] + totals["refill"]),
                        ("Final_Prize", totals["final"]),
                        ("Prizes_Won", totals["begin"] + totals["refill"] - totals["final"]),
                    ):
                        column = machine_headers.get(name)
                        if column:
                            machine_sheet.cell(row_index, column).value = value
        database._save_atomic(workbook)
    finally:
        workbook.close()


def _refresh_refill_report(path_value: Any, records: list[dict[str, Any]]) -> None:
    path = Path(str(path_value or ""))
    if not path.exists():
        return
    workbook = load_workbook(path)
    try:
        if "Product Detail" in workbook.sheetnames:
            del workbook["Product Detail"]
        sheet = workbook.create_sheet("Product Detail")
        headers = ["Machine ID", "Machine Name", "Barcode", "Begin Qty", "Refill Qty", "Final Qty", "Qty Used"]
        sheet.append(headers)
        _style_header(sheet)
        for item in records:
            sheet.append(
                [
                    item["machine_id"],
                    item["machine_name"],
                    item["barcode"],
                    item["begin_qty"],
                    item["refill_qty"],
                    item["final_qty"],
                    item["begin_qty"] + item["refill_qty"] - item["final_qty"],
                ]
            )
        for index, width in enumerate([16, 26, 24, 14, 14, 14, 14], start=1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
        workbook.save(path)
    finally:
        workbook.close()


def _take_api_endpoint(api, path: str, method: str):
    endpoint = None
    retained = []
    for route in api.router.routes:
        route_path = getattr(route, "path", "")
        methods = set(getattr(route, "methods", set()) or set())
        if route_path == path and method in methods:
            endpoint = getattr(route, "endpoint", endpoint)
            continue
        retained.append(route)
    api.router.routes[:] = retained
    if endpoint is None:
        raise RuntimeError(f"Cannot extend missing API route: {method} {path}")
    return endpoint


def _install_refill_v2138(namespace: dict[str, Any]) -> None:
    import database as database_module

    database_class = namespace["ExcelDatabase"]
    api = namespace["api"]
    db: Callable[[], Any] = namespace["db"]

    # v2.1.39: keep the established Product_Closing header order intact and append
    # refill columns. This lets existing v2.1.37/v2.1.38 workbooks migrate safely
    # before strict schema validation runs.
    for column in ("Refill_Qty", "Refill_History_JSON"):
        if column not in PRODUCT_CLOSING_HEADERS:
            PRODUCT_CLOSING_HEADERS.append(column)
    database_module.SCHEMA["Product_Closing"] = list(PRODUCT_CLOSING_HEADERS)

    current_create = database_class.create_default_database
    current_migrate = database_class._migrate_schema

    @staticmethod
    def create_default_database_v2138(path: Path) -> None:
        current_create(path)
        _upgrade_refill_workbook(Path(path))

    def migrate_schema_v2138(self) -> None:
        current_migrate(self)
        _upgrade_refill_workbook(self.path)

    database_class.create_default_database = create_default_database_v2138
    database_class._migrate_schema = migrate_schema_v2138

    original_save = _take_api_endpoint(api, "/api/closings/save", "POST")
    original_calculate = _take_api_endpoint(api, "/api/calculate", "POST")
    original_new_closing = _take_api_endpoint(api, "/api/new-closing", "GET")
    original_pdf = _take_api_endpoint(api, "/api/reports/pdf", "POST")

    def calculate_v2138(payload: dict[str, Any]):
        transformed, _records = _payload_with_refill_available(payload)
        return original_calculate(transformed)

    def new_closing_v2138(report_date: str):
        result = original_new_closing(report_date)
        for machine in result.get("machines", []):
            for product in machine.get("products", []):
                product["refill_qty"] = 0
                product["refill_history"] = []
        return result

    def save_closing_v2138(payload: dict[str, Any]):
        transformed, records = _payload_with_refill_available(payload)
        result = original_save(transformed)
        closing_id = str(result.get("closing_id") or payload.get("closing_id") or "")
        database = db()
        _patch_saved_refills(database, closing_id, records)
        if str(result.get("workflow_status") or "").casefold() == "finalized":
            _refresh_refill_report(result.get("report_path"), records)
        return result

    def export_pdf_v2138(payload: dict[str, Any]):
        transformed, _records = _payload_with_refill_available(payload)
        return original_pdf(transformed)

    api.add_api_route("/api/calculate", calculate_v2138, methods=["POST"])
    api.add_api_route("/api/new-closing", new_closing_v2138, methods=["GET"])
    api.add_api_route("/api/closings/save", save_closing_v2138, methods=["POST"])
    api.add_api_route("/api/reports/pdf", export_pdf_v2138, methods=["POST"])


_install_v217_before_refill = install_v217

def install_v217(namespace: dict[str, Any]) -> None:
    _install_v217_before_refill(namespace)
    _install_refill_v2138(namespace)

# v2.1.49 - operational Closing History, audit detail and protected finalized deletion.
def _install_history_v2149(namespace: dict[str, Any]) -> None:
    api = namespace["api"]
    db: Callable[[], Any] = namespace["db"]
    serialise = namespace["serialise"]
    REPORTS_DIR: Path = namespace["REPORTS_DIR"]
    export_daily_closing = namespace["export_daily_closing"]

    def closing_key(row: dict[str, Any]) -> tuple[str, str, str]:
        return (
            str(row.get("Report_Date") or "")[:10],
            str(row.get("Finalized_At") or row.get("Created_At") or row.get("Updated_At") or ""),
            str(row.get("Closing_ID") or ""),
        )

    def history_snapshot(limit: int = Query(1000, ge=1, le=5000)) -> dict[str, Any]:
        database = db()
        records = [dict(row) for row in database.list_closings(limit) if str(row.get("Workflow_Status") or "").casefold() == "finalized"]
        product_rows: list[dict[str, Any]] = []
        machine_rows: list[dict[str, Any]] = []
        with database._workbook(read_only=True) as workbook:
            if "Product_Closing" in workbook.sheetnames:
                product_rows = [dict(row) for row in database._sheet_rows(workbook["Product_Closing"])]
            if "Machine_Closing" in workbook.sheetnames:
                machine_rows = [dict(row) for row in database._sheet_rows(workbook["Machine_Closing"])]

        products_by_closing: dict[str, list[dict[str, Any]]] = {}
        for product in product_rows:
            products_by_closing.setdefault(str(product.get("Closing_ID") or ""), []).append(product)
        machines_by_closing: dict[str, set[str]] = {}
        for machine in machine_rows:
            cid = str(machine.get("Closing_ID") or "")
            mid = str(machine.get("Machine_ID") or "").strip()
            if mid:
                machines_by_closing.setdefault(cid, set()).add(mid)

        all_machines: set[str] = set()
        all_staff: set[str] = set()
        enriched: list[dict[str, Any]] = []
        for row in records:
            cid = str(row.get("Closing_ID") or "")
            products = products_by_closing.get(cid, [])
            refill_total = 0
            barcodes: set[str] = set()
            product_machines: set[str] = set()
            for product in products:
                refill_total += _refill_total(product)
                barcode = str(product.get("Barcode") or "").strip()
                machine_id = str(product.get("Machine_ID") or "").strip()
                if barcode:
                    barcodes.add(barcode)
                if machine_id:
                    product_machines.add(machine_id)
            machine_ids = sorted(machines_by_closing.get(cid, set()) | product_machines)
            all_machines.update(machine_ids)
            closed_by = str(row.get("Closed_By") or "").strip()
            if closed_by:
                all_staff.add(closed_by)
            item = dict(row)
            item["Refill_Qty"] = refill_total
            item["Machine_Ids"] = machine_ids
            item["Barcodes"] = sorted(barcodes)
            enriched.append(item)
        enriched.sort(key=closing_key, reverse=True)
        return serialise({"records": enriched, "machines": sorted(all_machines), "staff": sorted(all_staff, key=str.casefold)})

    def delete_finalized_closing(closing_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        database = db()
        target_id = str(closing_id or "").strip()
        reason = str(payload.get("reason") or "").strip()
        if len(reason) < 3:
            raise ValueError("Void Reason is required.")
        all_rows = database.list_closings(5000)
        target = next((dict(row) for row in all_rows if str(row.get("Closing_ID") or "") == target_id), None)
        if target is None:
            raise HTTPException(status_code=404, detail="Closing was not found.")
        if str(target.get("Workflow_Status") or "").casefold() != "finalized":
            raise ValueError("Only finalized closings can be voided from Closing History.")

        target_key = closing_key(target)
        newer_finalized = [dict(row) for row in all_rows if str(row.get("Workflow_Status") or "").casefold() == "finalized" and closing_key(dict(row)) > target_key]
        if newer_finalized:
            newest = sorted(newer_finalized, key=closing_key)[0]
            raise ValueError(f"Cannot void {target_id}. Newer finalized bill {newest.get('Closing_ID')} depends on it. Void newer finalized bills first.")
        dependent_drafts = [dict(row) for row in all_rows if str(row.get("Workflow_Status") or "").casefold() == "draft" and str(row.get("Report_Date") or "")[:10] >= str(target.get("Report_Date") or "")[:10]]
        draft_ids = [str(row.get("Closing_ID") or "").strip() for row in dependent_drafts if str(row.get("Closing_ID") or "").strip()]
        ids_to_remove = {target_id, *draft_ids}

        workbook = load_workbook(database.path)
        try:
            deleted_counts: dict[str, int] = {}
            for sheet_name in ("Daily_Closing", "Machine_Closing", "Product_Closing"):
                if sheet_name not in workbook.sheetnames:
                    continue
                sheet = workbook[sheet_name]
                headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
                closing_col = headers.get("Closing_ID")
                if not closing_col:
                    continue
                rows_to_delete = [row_index for row_index in range(2, sheet.max_row + 1) if str(sheet.cell(row_index, closing_col).value or "") in ids_to_remove]
                for row_index in reversed(rows_to_delete):
                    sheet.delete_rows(row_index, 1)
                deleted_counts[sheet_name] = len(rows_to_delete)
            for draft in dependent_drafts:
                draft_id = str(draft.get("Closing_ID") or "")
                database._append_audit(
                    workbook,
                    "VOID_DRAFT",
                    "Closing",
                    draft_id,
                    "Administrator",
                    {
                        "reason": f"Automatically removed because finalized bill {target_id} was voided.",
                        "parent_closing_id": target_id,
                        "parent_void_reason": reason,
                        "report_date": draft.get("Report_Date"),
                    },
                )

            database._append_audit(
                workbook,
                "VOID",
                "Closing",
                target_id,
                "Administrator",
                {
                    "reason": reason,
                    "report_date": target.get("Report_Date"),
                    "total_sales": target.get("Total_Sales"),
                    "coins_used": target.get("Machine_Coins_Used"),
                    "products_used": target.get("Total_Prizes_Won"),
                    "closed_by": target.get("Closed_By"),
                    "verified_by": target.get("Verified_By"),
                    "dependent_drafts_removed": draft_ids,
                    "deleted_counts": deleted_counts,
                },
            )
            database._save_atomic(workbook)
        finally:
            workbook.close()

        report_deleted = False
        report_delete_errors: list[str] = []
        for report_path in REPORTS_DIR.glob(f"{target_id}.*"):
            if not report_path.is_file():
                continue
            try:
                report_path.unlink(missing_ok=True)
                report_deleted = True
            except OSError as exc:
                report_delete_errors.append(str(exc))
        return serialise({"ok": True, "closing_id": target_id, "voided": True, "dependent_drafts_removed": draft_ids, "report_deleted": report_deleted, "report_delete_errors": report_delete_errors})

    def export_history_report(closing_id: str) -> dict[str, Any]:
        database = db()
        try:
            header, machines = database.get_closing(closing_id)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        if str(header.get("Workflow_Status") or "").casefold() != "finalized":
            raise ValueError("Only finalized closings have an export report.")
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        report_path = REPORTS_DIR / f"{closing_id}.xlsx"
        if not report_path.exists():
            export_daily_closing(header, machines, report_path)
            products = database.get_closing_products(closing_id)
            _append_product_detail_sheet(report_path, products)
            refill_records = [
                {
                    "machine_id": str(product.get("Machine_ID") or ""),
                    "machine_name": str(product.get("Machine_Name") or ""),
                    "barcode": str(product.get("Barcode") or ""),
                    "begin_qty": max(0, _to_int(product.get("Begin_Qty"))),
                    "refill_qty": _refill_total(product),
                    "final_qty": max(0, _to_int(product.get("Final_Qty"))),
                }
                for product in products
            ]
            _refresh_refill_report(report_path, refill_records)
        try:
            if os.name == "nt" and hasattr(os, "startfile"):
                os.startfile(str(report_path))
            else:
                import webbrowser
                webbrowser.open(report_path.resolve().as_uri())
        except OSError as exc:
            raise ValueError(f"The Excel report could not be opened: {exc}") from exc
        return {"ok": True, "report_name": report_path.name, "report_path": str(report_path)}

    api.add_api_route("/api/history", history_snapshot, methods=["GET"])
    api.add_api_route("/api/closings/{closing_id}", delete_finalized_closing, methods=["DELETE"])
    api.add_api_route("/api/history/{closing_id}/export", export_history_report, methods=["POST"])


_install_v217_before_history_v2149 = install_v217

def install_v217(namespace: dict[str, Any]) -> None:
    _install_v217_before_history_v2149(namespace)
    _install_history_v2149(namespace)


# v2.1.63 - signed adjustments and same-day carry forward
def _normalise_refill_history(value: Any) -> list[dict[str, Any]]:
    if isinstance(value,list): raw=value
    else:
        text=str(value or '').strip()
        if not text:return []
        try:raw=json.loads(text)
        except (json.JSONDecodeError,TypeError):return []
    if not isinstance(raw,list):return []
    out=[]
    for item in raw:
        if not isinstance(item,dict):continue
        qty=_to_int(item.get('qty'))
        if qty==0:continue
        out.append({'qty':qty,'at':str(item.get('at') or item.get('timestamp') or '')[:40],'by':str(item.get('by') or item.get('staff') or '')[:100]})
    return out[-100:]
def _refill_total(product:dict[str,Any])->int:
    h=_normalise_refill_history(product.get('refill_history') or product.get('Refill_History_JSON'))
    if h:return sum(_to_int(x.get('qty')) for x in h)
    return _to_int(product.get('refill_qty') if product.get('refill_qty') is not None else product.get('Refill_Qty'))
def _refill_records(payload_machines:list[dict[str,Any]])->list[dict[str,Any]]:
    records=[]
    for machine in payload_machines:
        mid=str(machine.get('machine_id') or '').strip();mname=str(machine.get('machine_name') or '').strip()
        for idx,product in enumerate(list(machine.get('products') or [])):
            barcode=str(product.get('barcode') or '').strip();pid=_safe_product_id(product.get('product_id'),mid,barcode,idx);h=_normalise_refill_history(product.get('refill_history') or product.get('Refill_History_JSON'));rq=sum(_to_int(x.get('qty')) for x in h) if h else _to_int(product.get('refill_qty') if product.get('refill_qty') is not None else product.get('Refill_Qty'));records.append({'machine_id':mid,'machine_name':mname,'product_id':pid,'barcode':barcode,'begin_qty':max(0,_to_int(product.get('begin_qty'))),'refill_qty':rq,'final_qty':max(0,_to_int(product.get('final_qty'))),'refill_history':h})
    return records
def _payload_with_refill_available(payload:dict[str,Any]):
    copied=dict(payload);originals=_refill_records(list(payload.get('machines') or []));byid={(x['machine_id'],x['product_id'].casefold()):x for x in originals};bybc={(x['machine_id'],x['barcode'].casefold()):x for x in originals if x['barcode']};cms=[]
    for machine in list(payload.get('machines') or []):
        mc=dict(machine);mid=str(machine.get('machine_id') or '').strip();ps=[]
        for idx,product in enumerate(list(machine.get('products') or [])):
            pc=dict(product);barcode=str(product.get('barcode') or '').strip();pid=_safe_product_id(product.get('product_id'),mid,barcode,idx);orig=byid.get((mid,pid.casefold())) or bybc.get((mid,barcode.casefold()));rq=int(orig['refill_qty']) if orig else _refill_total(pc);available=max(0,_to_int(product.get('begin_qty')))+rq
            if available<0:raise ValueError(f"{machine.get('machine_name') or mid} · {barcode or 'product'}: stock adjustment cannot make available quantity negative.")
            pc['begin_qty']=available;ps.append(pc)
        mc['products']=ps;cms.append(mc)
    copied['machines']=cms;return copied,originals
_install_v217_before_v2163=install_v217
def install_v217(namespace:dict[str,Any])->None:
    _install_v217_before_v2163(namespace);cls=namespace['ExcelDatabase']
    def previous_product_states_v2163(self,machine_ids,before_date=None):
        req={str(x) for x in machine_ids};finalized={}
        for row in self.list_closings(5000):
            if str(row.get('Workflow_Status') or '')!='Finalized':continue
            rd=str(row.get('Report_Date') or '')[:10]
            if before_date and rd>before_date.isoformat():continue
            finalized[str(row.get('Closing_ID') or '')]=(rd,str(row.get('Finalized_At') or row.get('Updated_At') or row.get('Created_At') or ''),str(row.get('Closing_ID') or ''))
        with self._workbook(read_only=True) as wb:rows=self._sheet_rows(wb['Product_Closing'])
        latest={}
        for row in rows:
            mid=str(row.get('Machine_ID') or '')
            if mid not in req:continue
            key=finalized.get(str(row.get('Closing_ID') or ''))
            if key is None:continue
            pk=str(row.get('Product_ID') or row.get('Barcode') or '').casefold()
            if not pk:continue
            cur=latest.get((mid,pk))
            if cur is None or key>cur[0]:latest[(mid,pk)]=(key,dict(row))
        result={}
        for (mid,pk),value in latest.items():
            result.setdefault(mid,{})[pk]=value[1];bc=str(value[1].get('Barcode') or '').casefold()
            if bc:result[mid][bc]=value[1]
        return result
    cls.previous_product_states=previous_product_states_v2163
