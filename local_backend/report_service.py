from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "102A43"
NAVY_2 = "173F5F"
GOLD = "D4A84F"
GREEN = "168A5B"
RED = "C0392B"
LIGHT = "F4F7FA"
MID = "D9E2EC"
WHITE = "FFFFFF"
TEXT = "243B53"


def _title(sheet, text: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=end_column)
    cell = sheet.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=end_column):
        for item in row:
            item.fill = PatternFill("solid", fgColor=NAVY)


def _section(sheet, row: int, text: str, start_column: int, end_column: int) -> None:
    sheet.merge_cells(start_row=row, start_column=start_column, end_row=row, end_column=end_column)
    cell = sheet.cell(row, start_column, text)
    cell.fill = PatternFill("solid", fgColor=NAVY_2)
    cell.font = Font(name="Aptos", bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[row].height = 24


def _style_header(row) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="FFFFFF"))


def export_daily_closing(
    header: dict[str, Any],
    machines: list[dict[str, Any]],
    destination: Path,
) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Closing Summary"
    detail = workbook.create_sheet("Machine Detail")

    _title(summary, "JOLI POLI  •  DAILY CLAW MACHINE CLOSING", 8)
    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A7"

    labels = [
        ("Closing ID", header.get("Closing_ID")),
        ("Report Date", header.get("Report_Date")),
        ("Outlet", header.get("Outlet")),
        ("Workflow", header.get("Workflow_Status")),
        ("Closed By", header.get("Closed_By")),
        ("Verified By", header.get("Verified_By")),
    ]
    for index, (label, value) in enumerate(labels):
        row = 3 + index // 3
        column = 1 + (index % 3) * 2
        summary.cell(row, column, label).font = Font(bold=True, color=TEXT)
        summary.cell(row, column + 1, value).fill = PatternFill("solid", fgColor="FFF4D6")
        summary.cell(row, column + 1).alignment = Alignment(horizontal="center")

    _section(summary, 6, "FINANCIAL & COIN CONTROL", 1, 8)
    financial_headers = [
        "Cash Sales",
        "ABA Sales",
        "Adjustment",
        "Total Sales",
        "Coins Dispensed",
        "Machine Coins",
        "Variance",
        "Status",
    ]
    financial_values = [
        header.get("Cash_Sales"),
        header.get("ABA_Sales"),
        header.get("Adjustment"),
        header.get("Total_Sales"),
        header.get("Coins_Dispensed"),
        header.get("Machine_Coins_Used"),
        header.get("Coin_Variance"),
        header.get("Closing_Status"),
    ]
    for column, value in enumerate(financial_headers, 1):
        summary.cell(7, column, value)
    _style_header(summary[7])
    for column, value in enumerate(financial_values, 1):
        cell = summary.cell(8, column, value)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor=LIGHT)
    for column in range(1, 5):
        summary.cell(8, column).number_format = '$#,##0.00;[Red]-$#,##0.00'
    summary.cell(8, 7).font = Font(
        bold=True, color=GREEN if int(header.get("Coin_Variance") or 0) == 0 else RED
    )
    summary.cell(8, 8).font = Font(
        bold=True, color=GREEN if header.get("Closing_Status") == "Balanced" else RED
    )

    _section(summary, 10, "PERFORMANCE SUMMARY", 1, 8)
    perf_headers = [
        "Total Prizes",
        "Avg Coins / Prize",
        "Avg Revenue / Prize",
        "Avg Value / Coin",
        "Cash Txns",
        "ABA Txns",
        "Total Txns",
        "Finalized At",
    ]
    perf_values = [
        header.get("Total_Prizes_Won"),
        header.get("Average_Coins_Per_Prize"),
        header.get("Average_Revenue_Per_Prize"),
        header.get("Average_Sale_Value_Per_Coin"),
        header.get("Cash_Transactions"),
        header.get("ABA_Transactions"),
        header.get("Total_Transactions"),
        header.get("Finalized_At"),
    ]
    for column, value in enumerate(perf_headers, 1):
        summary.cell(11, column, value)
    _style_header(summary[11])
    for column, value in enumerate(perf_values, 1):
        summary.cell(12, column, value)
        summary.cell(12, column).alignment = Alignment(horizontal="center")
        summary.cell(12, column).fill = PatternFill("solid", fgColor=LIGHT)
    summary.cell(12, 3).number_format = '$#,##0.00'
    summary.cell(12, 4).number_format = '$0.0000'

    _section(summary, 14, "MACHINE TYPE SUMMARY", 1, 8)
    type_summary: dict[str, dict[str, float]] = defaultdict(
        lambda: {"machines": 0, "coins": 0, "prizes": 0, "sales": 0.0}
    )
    for machine in machines:
        bucket = type_summary[str(machine.get("Machine_Type") or "Unknown")]
        bucket["machines"] += 1
        bucket["coins"] += int(machine.get("Coins_Used") or 0)
        bucket["prizes"] += int(machine.get("Prizes_Won") or 0)
        bucket["sales"] += float(machine.get("Allocated_Sales") or 0)

    headers = [
        "Machine Type",
        "Machines",
        "Coins Used",
        "Prizes Won",
        "Coins / Prize",
        "Allocated Sales",
        "Avg Revenue / Prize",
        "Sales Share",
    ]
    for column, value in enumerate(headers, 1):
        summary.cell(15, column, value)
    _style_header(summary[15])
    total_sales = float(header.get("Total_Sales") or 0)
    row = 16
    for machine_type, values in type_summary.items():
        coins_per = values["coins"] / values["prizes"] if values["prizes"] else 0
        avg_revenue = values["sales"] / values["prizes"] if values["prizes"] else 0
        share = values["sales"] / total_sales if total_sales else 0
        values_row = [
            machine_type,
            values["machines"],
            values["coins"],
            values["prizes"],
            coins_per,
            values["sales"],
            avg_revenue,
            share,
        ]
        for column, value in enumerate(values_row, 1):
            summary.cell(row, column, value)
            summary.cell(row, column).fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else LIGHT)
            summary.cell(row, column).alignment = Alignment(horizontal="center")
        summary.cell(row, 6).number_format = '$#,##0.00'
        summary.cell(row, 7).number_format = '$#,##0.00'
        summary.cell(row, 8).number_format = '0.00%'
        row += 1

    notes_row = row + 2
    _section(summary, notes_row, "NOTES", 1, 8)
    summary.merge_cells(start_row=notes_row + 1, start_column=1, end_row=notes_row + 4, end_column=8)
    notes_cell = summary.cell(notes_row + 1, 1, header.get("Notes") or "")
    notes_cell.alignment = Alignment(vertical="top", wrap_text=True)
    notes_cell.fill = PatternFill("solid", fgColor="FFF4D6")

    for column, width in enumerate([22, 18, 18, 18, 18, 18, 20, 22], 1):
        summary.column_dimensions[get_column_letter(column)].width = width
    summary.print_area = f"A1:H{notes_row + 4}"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.sheet_properties.pageSetUpPr.fitToPage = True

    _title(detail, "MACHINE CLOSING DETAIL", 18)
    detail.sheet_view.showGridLines = False
    detail.freeze_panes = "A5"
    detail_headers = [
        "Machine ID",
        "Machine Name",
        "Type",
        "Capacity",
        "Begin Prize",
        "Refill",
        "Available",
        "Final Prize",
        "Prizes Won",
        "Refill Needed",
        "Begin Meter",
        "Final Meter",
        "Coins Used",
        "Coins / Prize",
        "Allocated Sales",
        "Avg Revenue / Prize",
        "Status",
        "Notes",
    ]
    for column, value in enumerate(detail_headers, 1):
        detail.cell(4, column, value)
    _style_header(detail[4])
    for row_idx, machine in enumerate(machines, 5):
        values = [
            machine.get("Machine_ID"),
            machine.get("Machine_Name"),
            machine.get("Machine_Type"),
            machine.get("Capacity"),
            machine.get("Begin_Prize"),
            machine.get("Refill_Prize"),
            machine.get("Available_Prize"),
            machine.get("Final_Prize"),
            machine.get("Prizes_Won"),
            machine.get("Refill_Needed"),
            machine.get("Begin_Coin_Meter"),
            machine.get("Final_Coin_Meter"),
            machine.get("Coins_Used"),
            machine.get("Coins_Per_Prize"),
            machine.get("Allocated_Sales"),
            machine.get("Average_Revenue_Per_Prize"),
            machine.get("Machine_Status"),
            machine.get("Notes"),
        ]
        for column, value in enumerate(values, 1):
            cell = detail.cell(row_idx, column, value)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_idx % 2 else LIGHT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        detail.cell(row_idx, 15).number_format = '$#,##0.00'
        detail.cell(row_idx, 16).number_format = '$#,##0.00'

    for column in range(1, 19):
        detail.column_dimensions[get_column_letter(column)].width = 15
    detail.column_dimensions["B"].width = 24
    detail.column_dimensions["C"].width = 18
    detail.column_dimensions["R"].width = 32
    detail.auto_filter.ref = f"A4:R{max(5, 4 + len(machines))}"
    detail.print_title_rows = "1:4"
    detail.page_setup.orientation = "landscape"
    detail.page_setup.fitToWidth = 1
    detail.sheet_properties.pageSetUpPr.fitToPage = True

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def export_monthly_summary(
    month_text: str,
    closings: list[dict[str, Any]],
    machine_rows: list[dict[str, Any]],
    destination: Path,
) -> Path:
    workbook = Workbook()
    daily = workbook.active
    daily.title = "Monthly Summary"
    machine_sheet = workbook.create_sheet("Machine Performance")

    _title(daily, f"JOLI POLI  •  MONTHLY CLAW REPORT  •  {month_text}", 10)
    daily.sheet_view.showGridLines = False
    daily.freeze_panes = "A5"
    headers = [
        "Date",
        "Closing ID",
        "Cash Sales",
        "ABA Sales",
        "Adjustment",
        "Total Sales",
        "Coins Used",
        "Prizes Won",
        "Coin Variance",
        "Verified By",
    ]
    for column, value in enumerate(headers, 1):
        daily.cell(4, column, value)
    _style_header(daily[4])
    for row_idx, closing in enumerate(sorted(closings, key=lambda item: str(item.get("Report_Date"))), 5):
        values = [
            closing.get("Report_Date"),
            closing.get("Closing_ID"),
            closing.get("Cash_Sales"),
            closing.get("ABA_Sales"),
            closing.get("Adjustment"),
            closing.get("Total_Sales"),
            closing.get("Machine_Coins_Used"),
            closing.get("Total_Prizes_Won"),
            closing.get("Coin_Variance"),
            closing.get("Verified_By"),
        ]
        for column, value in enumerate(values, 1):
            cell = daily.cell(row_idx, column, value)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_idx % 2 else LIGHT)
            cell.alignment = Alignment(horizontal="center")
        for column in range(3, 7):
            daily.cell(row_idx, column).number_format = '$#,##0.00'

    summary_row = 5 + len(closings) + 2
    daily.cell(summary_row, 1, "MONTH TOTAL").font = Font(bold=True, color=WHITE)
    daily.cell(summary_row, 1).fill = PatternFill("solid", fgColor=NAVY)
    for column in range(2, 11):
        daily.cell(summary_row, column).fill = PatternFill("solid", fgColor=NAVY)
        daily.cell(summary_row, column).font = Font(bold=True, color=WHITE)
    daily.cell(summary_row, 3, sum(float(row.get("Cash_Sales") or 0) for row in closings))
    daily.cell(summary_row, 4, sum(float(row.get("ABA_Sales") or 0) for row in closings))
    daily.cell(summary_row, 5, sum(float(row.get("Adjustment") or 0) for row in closings))
    daily.cell(summary_row, 6, sum(float(row.get("Total_Sales") or 0) for row in closings))
    daily.cell(summary_row, 7, sum(int(row.get("Machine_Coins_Used") or 0) for row in closings))
    daily.cell(summary_row, 8, sum(int(row.get("Total_Prizes_Won") or 0) for row in closings))
    daily.cell(summary_row, 9, sum(int(row.get("Coin_Variance") or 0) for row in closings))
    for column in range(3, 7):
        daily.cell(summary_row, column).number_format = '$#,##0.00'

    for column, width in enumerate([14, 20, 16, 16, 16, 16, 14, 14, 14, 20], 1):
        daily.column_dimensions[get_column_letter(column)].width = width

    _title(machine_sheet, f"MACHINE PERFORMANCE  •  {month_text}", 9)
    machine_headers = [
        "Machine ID",
        "Machine Name",
        "Type",
        "Closings",
        "Coins Used",
        "Prizes Won",
        "Coins / Prize",
        "Allocated Sales",
        "Avg Revenue / Prize",
    ]
    for column, value in enumerate(machine_headers, 1):
        machine_sheet.cell(4, column, value)
    _style_header(machine_sheet[4])

    aggregate: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"name": "", "type": "", "closings": 0, "coins": 0, "prizes": 0, "sales": 0.0}
    )
    for row in machine_rows:
        machine_id = str(row.get("Machine_ID") or "")
        item = aggregate[machine_id]
        item["name"] = row.get("Machine_Name")
        item["type"] = row.get("Machine_Type")
        item["closings"] += 1
        item["coins"] += int(row.get("Coins_Used") or 0)
        item["prizes"] += int(row.get("Prizes_Won") or 0)
        item["sales"] += float(row.get("Allocated_Sales") or 0)

    row_idx = 5
    for machine_id, item in sorted(aggregate.items()):
        coins_per = item["coins"] / item["prizes"] if item["prizes"] else 0
        avg_revenue = item["sales"] / item["prizes"] if item["prizes"] else 0
        values = [
            machine_id,
            item["name"],
            item["type"],
            item["closings"],
            item["coins"],
            item["prizes"],
            coins_per,
            item["sales"],
            avg_revenue,
        ]
        for column, value in enumerate(values, 1):
            machine_sheet.cell(row_idx, column, value)
            machine_sheet.cell(row_idx, column).fill = PatternFill(
                "solid", fgColor=WHITE if row_idx % 2 else LIGHT
            )
            machine_sheet.cell(row_idx, column).alignment = Alignment(horizontal="center")
        machine_sheet.cell(row_idx, 8).number_format = '$#,##0.00'
        machine_sheet.cell(row_idx, 9).number_format = '$#,##0.00'
        row_idx += 1

    if aggregate:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Allocated Sales by Machine"
        chart.y_axis.title = "Machine"
        chart.x_axis.title = "Sales"
        data = Reference(machine_sheet, min_col=8, min_row=4, max_row=row_idx - 1)
        categories = Reference(machine_sheet, min_col=2, min_row=5, max_row=row_idx - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 8
        chart.width = 17
        machine_sheet.add_chart(chart, "K4")

    for column, width in enumerate([14, 24, 18, 12, 14, 14, 16, 18, 20], 1):
        machine_sheet.column_dimensions[get_column_letter(column)].width = width
    machine_sheet.freeze_panes = "A5"
    machine_sheet.sheet_view.showGridLines = False

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination
