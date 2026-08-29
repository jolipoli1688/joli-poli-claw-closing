"""Offline-only migration preflight. It never opens a network connection or writes source data."""
from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
from openpyxl import load_workbook

REQUIRED_SHEETS = ("Settings", "Machine_Master", "Daily_Closing", "Machine_Closing")
def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
def values(ws, header: str) -> list[str]:
    headers = {str(cell.value or ""): index for index, cell in enumerate(next(ws.iter_rows(values_only=True)), start=0)}
    if header not in headers: raise ValueError(f"{ws.title} is missing {header}")
    return [str(row[headers[header]] or "").strip() for row in ws.iter_rows(values_only=True) if str(row[headers[header]] or "").strip()]
def duplicates(items: list[str]) -> list[str]:
    seen, result = set(), []
    for item in items:
        if item in seen and item not in result: result.append(item)
        seen.add(item)
    return result
def main() -> int:
    parser = argparse.ArgumentParser(); parser.add_argument("--source-workbook", type=Path, required=True); parser.add_argument("--dry-run", action="store_true", required=True); parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    if not args.source_workbook.is_file(): raise SystemExit("source workbook was not found")
    wb = load_workbook(args.source_workbook, read_only=True, data_only=False)
    try:
        missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
        if missing: raise SystemExit("missing required sheets: " + ", ".join(missing))
        duplicate_ids = {"machine_ids": duplicates(values(wb["Machine_Master"], "Machine_ID")), "closing_ids": duplicates(values(wb["Daily_Closing"], "Closing_ID"))}
        failures = {name: ids for name, ids in duplicate_ids.items() if ids}
        if failures: raise SystemExit("duplicate source identifiers: " + json.dumps(failures, sort_keys=True))
        report = {"mode":"dry-run","source_sha256":digest(args.source_workbook),"sheets":{name:max(0, wb[name].max_row-1) for name in wb.sheetnames},"duplicate_ids":duplicate_ids,"id_mapping_report":{"strategy":"preserve legacy Machine_ID and Closing_ID as unique migration keys","records":[]},"idempotency":"target unique migration keys reject duplicate imports","network_attempted":False,"writes_performed":False}
    finally: wb.close()
    text=json.dumps(report, indent=2, sort_keys=True)
    if args.report: args.report.write_text(text+"\n", encoding="utf-8")
    print(text); return 0
if __name__ == "__main__": raise SystemExit(main())
